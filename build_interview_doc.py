from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = r"C:\Users\ASUS\Desktop\AI Agent\品质检查项目组\AI品质检查项目面试问答手册.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_width(cell, width_inch):
    width = Inches(width_inch)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inch * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            set_cell_width(row.cells[idx], width)


def add_heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.color.rgb = RGBColor(31, 78, 121) if level == 1 else RGBColor(47, 84, 150)
    return p


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.name = "Microsoft YaHei"
        r1._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r1.font.size = Pt(10.5)
        text = text[len(bold_prefix):]
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10.5)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(10)


def add_question(doc, idx, title, why, answer, followups=None, traps=None):
    add_heading(doc, f"Q{idx}. {title}", level=2)
    add_para(doc, f"面试官真正想考察：{why}", bold_prefix="面试官真正想考察：")
    add_para(doc, "建议回答：", bold_prefix="建议回答：")
    for para in answer:
        add_para(doc, para)
    if followups:
        add_para(doc, "可能追问与应对：", bold_prefix="可能追问与应对：")
        add_bullets(doc, followups)
    if traps:
        add_para(doc, "不要这样答：", bold_prefix="不要这样答：")
        add_bullets(doc, traps)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.5)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("AI品质检查项目面试问答手册")
run.bold = True
run.font.size = Pt(22)
run.font.name = "Microsoft YaHei"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
run.font.color.rgb = RGBColor(31, 78, 121)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = subtitle.add_run("基于《物业品质检查多Agent系统》项目文件与贺星AI产品经理简历整理")
sr.font.size = Pt(11)
sr.font.name = "Microsoft YaHei"
sr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

rows = [
    ("适用场景", "AI产品经理面试、项目深挖、技术方案复盘"),
    ("核心项目", "物业品质检查多Agent系统"),
    ("重点能力", "Multi-Agent架构判断、模型选型、评分链路、报告生成、整改闭环、效果评估"),
    ("使用方式", "先背“建议回答”的主线，再用“可能追问”补充实现细节"),
]
for k, v in rows:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(f"{k}：")
    r1.bold = True
    r1.font.name = "Microsoft YaHei"
    r1._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r1.font.color.rgb = RGBColor(31, 78, 121)
    r1.font.size = Pt(10.5)
    r2 = p.add_run(v)
    r2.font.name = "Microsoft YaHei"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r2.font.size = Pt(10.5)

doc.add_paragraph()
add_heading(doc, "一、面试表达总原则", level=1)
add_para(doc, "这个项目最容易被追问的风险点是：简历里写了 Multi-Agent、LangGraph、Qwen、DeepSeek、RAG、效果评估、自进化，但如果回答停留在名词层面，面试官会判断你只是包装项目。回答时要把技术名词还原成业务决策：为什么需要、怎么落地、如何评估、失败怎么兜底。")
add_bullets(doc, [
    "先讲业务流程：任务分配 -> 现场检查 -> AI评分 -> 报告生成 -> 整改复核 -> 趋势分析。",
    "再讲拆分原则：不同环节的输入输出、模型能力、成本频次和风险不同，所以拆成专门 Agent。",
    "最后讲治理能力：结构化输出、置信度、人工复核、重试熔断、模型降级、日志与成本监控。",
])

add_heading(doc, "二、项目一句话版本", level=1)
add_para(doc, "我的 AI 品质检查系统不是一个简单的检查表线上化工具，而是把物业品质检查的线下 SOP、432 项标准、评分规则、问题照片、整改记录和管理分析，抽象成一个可追溯、可复核、可持续优化的 AI 工作流。系统采用 Vue3 + FastAPI + SQLite + LangGraph，前端分为移动作业端和 PC 管理端；AI 层按任务拆成评分、报告、整改复核和综合分析等能力，评分用 Qwen-plus，报告和分析用 DeepSeek，并通过数据库进行 Agent 间数据传递，保证流程可审计、可恢复、可复用。")

questions = [
    {
        "title": "请你用 2 分钟介绍这个 AI 品质检查项目。",
        "why": "你是否真的理解项目的业务价值，而不是只会讲技术栈。",
        "answer": [
            "这个项目的业务背景是物业品质检查原来高度依赖人工。检查人员要按 8 大模块、432 项检查标准逐项检查，记录问题和照片；主管再人工评分、汇总、写报告、跟整改。痛点是检查项多、评分口径不一致、报告撰写慢、整改闭环难追踪。",
            "我的产品方案是把它拆成一个端到端闭环：PC 端创建任务和分配模块，移动端现场检查和拍照取证，检查完成后触发 AI 评分，评分结果再进入报告生成和整改闭环，最后支持跨项目、跨时间、跨区域的综合分析。",
            "技术上我采用 Vue3 + FastAPI + SQLite 的前后端分离架构，AI 编排用 LangGraph。Agent 不是随便堆的，而是按业务责任拆分：数据采集是规则驱动，不调用 LLM；评分 Agent 用 Qwen-plus 做结构化评分；报告和综合分析用 DeepSeek 处理长文本总结；整改复核结合图片和整改说明进行审核。",
            "上线后的价值体现在三个层面：第一，检查标准 100% 结构化；第二，报告生成从 2-3 天压缩到 3 分钟内；第三，检查、评分、报告、整改、复核全链路线上化，闭环率从不足 40% 提升到 100%。"
        ],
        "followups": [
            "如果面试官要求更短：一句话说“我把线下品质检查 SOP 产品化成可评分、可复核、可追溯的 AI 工作流”。",
            "如果面试官问你的角色：强调你负责需求抽象、流程设计、Agent 分工、模型策略、评估指标和落地联调，而不是只写文档。"
        ]
    },
    {
        "title": "为什么选择多智能体框架，而不是一个大模型/一个 Agent 直接做完？",
        "why": "你是否能从业务复杂度、上下文、风险隔离和迭代效率解释架构选择。",
        "answer": [
            "我选择 Multi-Agent 不是因为它听起来先进，而是因为这个业务天然不是一个单一任务。品质检查至少包括数据采集、评分、报告、整改复核、综合分析五类动作，它们的输入、输出、评价标准和风险完全不同。",
            "如果用一个 Agent 做全部事情，系统 Prompt 会非常臃肿：它既要像评分专家，又要像报告专家，还要像整改审核员和经营分析师。这样会产生 Prompt 干扰，模型注意力被多个目标分散，评分口径和报告表达都容易不稳定。",
            "第二是上下文窗口问题。评分需要读取检查项、评分规则和问题描述；报告需要读取评分矩阵和问题分布；综合分析需要跨报告比较。如果所有上下文都塞给一个 Agent，Token 成本高，而且关键信息容易被噪声淹没。",
            "第三是故障隔离和迭代效率。评分逻辑有问题时，我只需要调整 Scoring Agent 的 Prompt、Few-shot 或回退规则；报告生成不受影响。整改复核要加入图片比对能力时，也不会污染评分链路。",
            "所以我的拆分原则是：业务责任单一、输入输出清晰、失败可隔离、评估指标不同，就拆成独立 Agent；但它们之间不做开放式聊天，而是通过数据库和 LangGraph 状态进行确定性流转。"
        ],
        "followups": [
            "如果问为什么不是单 Agent 加多个工具：工具调用适合开放式任务，这个项目是固定业务流程，确定性工作流更可靠。",
            "如果问 Multi-Agent 是否过度设计：回答“不是按概念拆，而是按业务环节拆；数据采集甚至不用 LLM，说明我不是为了 AI 而 AI”。"
        ],
        "traps": [
            "不要说“因为 Multi-Agent 更先进”。",
            "不要把 Agent 讲成互相聊天开会，这个项目强调确定性输入输出。"
        ]
    },
    {
        "title": "为什么用 LangGraph，而不是 CrewAI、AutoGen，或者自己写 for-loop？",
        "why": "你是否理解框架能力与业务流程的匹配关系。",
        "answer": [
            "我对这个问题的判断标准是：品质检查不是开放式头脑风暴，而是有明确状态、明确节点和明确失败处理的业务流程。所以我不需要 CrewAI 或 AutoGen 那种强调多角色对话的模式，我需要的是可控的状态机和工作流编排。",
            "LangGraph 的优势在于它把 Agent 流程显式建模成 StateGraph。比如评分链路是 collect_records -> fan-out(score_module x N) -> fan-in -> compute_total。报告链路是 collect_data -> fan-out(analyze_module x N) -> generate_report -> export_files -> save_report。每一步的状态、进度、错误都能被前端展示和后端追踪。",
            "如果自己写 for-loop 也能调用 8 次模型，但会缺少标准化状态管理、节点级可观测性、条件边和后续人工审核门控能力。项目后面加了 Orchestrator，把评分、审核门控、报告、整改追踪串起来，这种有状态流转用 LangGraph 更合适。",
            "所以我的选择不是“LangGraph 最好”，而是“在固定流程、需要状态追踪、需要并行 fan-out/fan-in、需要人工介入的业务场景下，LangGraph 比开放式 Agent 对话框架更稳”。"
        ],
        "followups": [
            "如果问是否会用其他框架：可以说“开放式协作/研究任务可以考虑 CrewAI、AutoGen；确定性业务流程优先 LangGraph 或 Temporal 类工作流”。",
            "如果问为什么不用 LangChain Agent：回答“我不需要模型自主决定工具链，而是希望产品流程可控、可审计”。"
        ]
    },
    {
        "title": "为什么评分用 Qwen-plus，报告和分析用 DeepSeek？",
        "why": "你是否能做成本、质量、稳定性之间的产品决策。",
        "answer": [
            "我的模型选型原则是按任务特征分配，而不是全系统用一个最贵或最强模型。评分是高频、结构化、规则约束强的任务；报告和综合分析是低频、长文本、综合推理和表达质量要求更高的任务。",
            "评分 Agent 选择 Qwen-plus，是因为评分输入主要是中文检查标准、评分规则、问题描述，输出是 JSON 结构化结果，包括 item_id、score、scoring_basis、improvement_suggestion、confidence。这个任务更看重中文规则理解、稳定结构化输出和成本控制。Qwen-plus 在这个场景里性价比更高。",
            "报告和分析选择 DeepSeek，是因为它要根据多个模块评分、问题分布、严重程度和建议生成专业报告，还要做跨项目、跨时间、跨区域对比分析。这类任务需要更强的长文本组织、综合归纳和业务表达能力，所以用 DeepSeek 更合适。",
            "另外我不是只做模型指定，还配了生产保障：评分侧有低温度、JSON 解析、重试、熔断、fallback 到 qwen-turbo、默认规则回退；报告侧有 DeepSeek 的重试和降级。这样模型选型不是单点选择，而是一套稳定性方案。"
        ],
        "followups": [
            "如果问为什么不用 DeepSeek 评分：回答“可以用，但评分是高频结构化任务，用 DeepSeek 成本不一定划算，且收益不明显”。",
            "如果问为什么不用 Qwen 生成报告：回答“可以作为降级，但正式报告更看重综合表达和长文本质量，DeepSeek 更适合主模型”。"
        ]
    },
    {
        "title": "你是怎么做意图识别的？",
        "why": "面试官在确认你是否把聊天机器人概念硬套到业务系统里。",
        "answer": [
            "这个问题我会先澄清：我的项目不是开放式聊天机器人，所以没有把所有请求都交给 LLM 做自由文本意图分类。品质检查是强流程产品，主意图识别更多是确定性路由：根据用户角色、页面入口、任务状态、按钮动作、分析模式参数来判断下一步执行哪个 Agent 或 API。",
            "例如用户在移动端完成检查后点击“开始评分”，这不是让模型猜意图，而是后端明确进入 Scoring Agent；评分完成后 Orchestrator 根据状态进入报告生成；综合分析页面用户选择 cross_project、cross_time 或 all_projects，后端先做参数校验，再进入对应分析链路。",
            "LLM 层也有一类更窄的“输入意图识别”：评分 Prompt 会判断问题描述是否与检查项相关。如果用户随便填“你好”“测试”“我试试”等无意义内容，模型不能正常评分，而是将 score 设为 3，confidence 设为 0.3 以下，并在 scoring_basis 中标注“疑似随意填写，建议人工复核”。",
            "所以我的设计是：业务主流程用规则和状态机保证确定性；只有在问题描述语义判断这种非结构化场景，才让模型参与识别，并用置信度和人工复核兜底。"
        ],
        "followups": [
            "如果面试官追问“有没有 intent classifier”：可以说“没有做通用分类器，因为需求不是开放式对话；如果未来加自然语言入口，我会用规则优先 + 轻量分类模型 + 置信度阈值”。",
            "如果追问具体字段：回答“综合分析模式用 mode 字段；评分流程用 task_id、record_id、module_name、status；异常输入通过 Prompt 规则和 confidence 标记”。"
        ],
        "traps": [
            "不要硬说用了复杂 NLP 意图识别模型。",
            "不要把前端路由、后端状态和 LLM 异常输入判断混成一件事。"
        ]
    },
    {
        "title": "Agent 之间是怎么通信的？为什么不用 Agent 之间直接调用？",
        "why": "考察你对可靠性、解耦和审计的理解。",
        "answer": [
            "系统里 Agent 之间不是互相发消息聊天，而是通过数据库作为共享数据层。比如评分 Agent 将每个检查项的 score、weighted_score、scoring_basis、improvement_suggestion 写入 scoring_results；报告 Agent 再从数据库读取评分结果和问题记录生成报告。",
            "这样做的核心原因是解耦和可审计。业务流程虽然是线性的，但运行时不应该强耦合。评分完成后，即使报告服务暂时失败，评分结果仍然在数据库里，可以重试报告生成；如果报告结果有问题，也能追溯到评分结果、检查记录和原始问题。",
            "直接 Agent-to-Agent 调用会让上游依赖下游的实时可用性，一旦下游超时，上游也失败。而数据库共享让每个 Agent 是生产者或消费者，天然支持重试、恢复、复用和审计。"
        ],
        "followups": [
            "如果问缺点：数据库通信会增加持久化设计成本，但对企业业务流程来说，审计和恢复价值更高。",
            "如果问是否实时：进度通过内存缓存、数据库双写、SSE 事件推送给前端。"
        ]
    },
    {
        "title": "评分 Agent 的完整流程是什么？",
        "why": "确认你是否知道核心 AI 链路，不只是知道“用 Qwen 打分”。",
        "answer": [
            "评分流程分三步。第一步 collect_records，从数据库查找当前检查任务下已完成的 InspectionRecord，并过滤已经评分过的模块。第二步 fan-out，对每个模块并行进入 score_module。第三步 compute_total，计算模块百分制得分和项目总分。",
            "在单模块评分里，系统先读取该模块模板检查项，再把数据库里的 Issue 按 item_id 分组。没有问题的合格项直接给 5 分，不进入 LLM；有问题的检查项才进入 Qwen-plus 批量评分。这样能减少 Token 和调用次数，也避免模型对合格项产生不必要波动。",
            "LLM 输出后会做多策略匹配：先按 item_id 精确匹配，再处理空格差异、序号映射、item_name 匹配，防止模型返回的 ID 格式小偏差导致结果丢失。最后保存 score、weight、weighted_score、评分依据和整改建议到 scoring_results。",
            "如果 AI 失败，系统不是直接崩溃，而是有重试、降级模型和规则回退评分。规则回退会按严重、一般、轻微问题数量扣分，而不是简单给一个固定中间分。"
        ],
        "followups": [
            "如果问为什么合格项不进 LLM：回答“合格项按规则给满分，降低成本并减少不确定性”。",
            "如果问批量粒度：回答“先按模块，再对有问题项按批次处理，避免单次 Prompt 过长或超时”。"
        ]
    },
    {
        "title": "评分公式是怎么设计的？",
        "why": "考察产品经理是否理解业务指标，而不是把评分交给模型黑箱。",
        "answer": [
            "评分不是完全由模型给一个项目总分。模型只负责根据检查项规则和问题描述输出单项评价得分，范围是 0-5 分。后续加权计算由后端规则完成。",
            "单项加权得分 = 评价得分 x 检查项权重。模块原始得分 = 模块内所有单项加权得分求和。因为每个模块权重总和为 0.2，模块满分是 5 x 0.2 = 1.0，再乘以 100 转成模块百分制得分。",
            "项目总分不是简单平均，而是按模块重要性加权。比如客户服务、安全管理、环境管理、机电运维、设施维护各占 15%，EHS 和综合管理各占 10%，财务管理占 5%。这样既保留检查项粒度，又符合管理层对不同模块重要性的判断。",
            "这个设计的关键是把模型和规则分工清楚：模型做语义判断和评分依据生成，系统做确定性公式计算，保证总分可解释、可复核。"
        ],
        "followups": [
            "如果问跳过项怎么处理：当前设计里 skipped 默认给 5 分，但真实上线需要限制跳过权限和要求填写跳过原因。",
            "如果问权重从哪里来：来自 Excel 检查表模板，通过字段映射读取。"
        ]
    },
    {
        "title": "你的 Prompt 是怎么设计的？",
        "why": "考察你是否理解 Prompt 工程和结构化输出。",
        "answer": [
            "评分 Prompt 不是简单一句“请打分”，它包括角色定义、评分体系、扣分原则、异常输入识别、输出格式、Few-shot 示例和历史记忆上下文。",
            "角色定义让模型稳定在“物业品质检查评分专家”；评分体系定义 0-5 分含义；扣分原则明确问题严重程度、数量、覆盖面、整改难度；输出格式强制 item_id、score、scoring_basis、improvement_suggestion、confidence 五个字段。",
            "Few-shot 示例覆盖 8 大模块，比如客户服务、安全管理、EHS、环境管理、机电运维、设施维护、综合管理、财务管理，帮助模型学会不同模块的扣分口径。历史记忆上下文则来自短记忆、长记忆和相似评分案例检索，用于减少同类问题前后评分不一致。",
            "报告 Prompt 则不是逐项评分，而是面向模块分析和综合报告：输入模块得分、评分依据、主要问题和建议，让 DeepSeek 输出模块评价、主要问题、整改建议和综合报告。"
        ],
        "followups": [
            "如果问 Prompt 版本管理：可以说“当前已有 prompt_manager 读取评分系统 Prompt，后续会做版本化、A/B 测试和回滚”。",
            "如果问如何防 JSON 失败：回答“强制 JSON 输出、低温度、解析代码提取 JSON array、失败重试和默认回退”。"
        ]
    },
    {
        "title": "你提到 RAG/记忆注入，它在评分里具体怎么用？",
        "why": "区分你是否真的理解 RAG，而不是把所有上下文都叫 RAG。",
        "answer": [
            "评分里的 RAG/记忆不是为了回答知识库问题，而是为了提升同类问题评分一致性。系统会在评分前构建 memory_context，包含三类信息：本次任务的短记忆、跨项目的长记忆、历史相似评分案例。",
            "短记忆解决同一次检查中前后模块的上下文一致性，比如某项目本次检查整体问题较多，模型可以参考前面模块的评分风格。长记忆解决跨项目经验复用，比如某类设施问题在历史中经常被判为一般或严重。相似案例检索则给模型提供参考评分记录。",
            "这样做的价值是把人工修正和历史评分经验沉淀下来，下一次评分时注入 Prompt，而不是每次都让模型从零判断。它本质上是一个持续改进机制。"
        ],
        "followups": [
            "如果问风险：历史案例可能带来偏差，所以需要人工审核、案例质量筛选和回滚机制。",
            "如果问和法律知识库 RAG 的区别：法律 RAG 是召回资料回答问题；评分 RAG 是召回历史案例校准打分口径。"
        ]
    },
    {
        "title": "AI 评分不准怎么办？",
        "why": "考察你是否有 AI 产品的质量治理思路。",
        "answer": [
            "我会把 AI 评分不准拆成三类：输入问题、模型输出问题、评分口径问题。输入问题包括检查员描述不清、随意填写、缺照片；输出问题包括 JSON 解析失败、item_id 匹配失败、置信度低；评分口径问题包括同类问题前后扣分不一致。",
            "对应策略也分三层。第一层是输入约束和异常识别，问题描述无效时 confidence 低于 0.3 并建议人工复核。第二层是结构化输出治理，强制 JSON、低温度、结果校验、重试、降级、规则回退。第三层是持续优化，把人工修正记录沉淀为纠偏案例和 Prompt 指令，下次评分时注入。",
            "产品上我不会承诺 AI 100% 正确，而是设计“AI 为主、人工兜底”的混合决策。高置信度结果自动通过，低置信度或严重问题进入人工复核；主管可以修改分数，修改原因再进入后续优化闭环。"
        ],
        "followups": [
            "如果问指标：一致率、人工修改率、低置信度占比、严重问题漏判率、JSON 解析成功率、平均耗时、Token 成本。",
            "如果问上线门槛：关键模块抽样人工复核，严重问题宁可多召回也不能漏判。"
        ]
    },
    {
        "title": "你怎么设计置信度和人工复核？",
        "why": "考察 AI 风险控制能力。",
        "answer": [
            "评分结果里除了 score，还有 confidence。置信度不是给用户看的装饰字段，而是驱动人工复核策略的信号。比如问题描述充分、评分规则明确、场景典型，置信度可以在 0.9 以上；描述模糊或场景少见时在 0.5-0.7；无意义输入或与检查项无关时低于 0.3。",
            "人工复核策略可以按三类触发：第一，confidence 低于阈值；第二，问题严重程度为“严重”；第三，AI 分数与历史同类案例偏差过大。复核结果包括是否接受 AI 分、人工修正分、修正原因。",
            "这些人工修正不是只存在一次，而是进入偏差检测和规则提炼流程，形成“偏差检测 -> 规则提炼 -> Prompt 注入 -> 效果验证 -> 回滚”的闭环。这样 AI 能力会随着业务使用持续改进。"
        ],
        "followups": [
            "如果问阈值怎么定：先用历史样本离线评估，再按业务风险调整，严重问题阈值更严格。",
            "如果问人工成本：只复核低置信度和高风险样本，不是全量复核。"
        ]
    },
    {
        "title": "报告生成 Agent 为什么也要拆出来？",
        "why": "考察你是否理解评分和报告生成的认知任务差异。",
        "answer": [
            "评分和报告是不同类型的任务。评分关注单项规则、扣分逻辑和结构化输出；报告关注模块归纳、问题排序、管理建议和专业表达。如果放在一个 Agent 里，评分 Prompt 和报告 Prompt 会互相干扰。",
            "报告 Agent 的流程是先从数据库收集评分结果、问题列表和模块得分，再对每个已评分模块并行做模块分析，最后汇总成综合报告并导出 Word/PDF。DeepSeek 在这里负责长文本分析和报告表达，而不是重新计算分数。",
            "这样拆分后，报告的可信度来自评分结果和原始问题的可追溯性。主管如果质疑报告结论，可以追溯到模块得分、检查项评分依据和问题照片。"
        ],
        "followups": [
            "如果问报告是否会幻觉：回答“报告输入来自结构化评分和问题数据，输出模板限定章节，关键数字由系统计算，不让模型自由编造”。",
            "如果问报告结构：检查概况、总体得分、模块分析、重点问题清单、整改建议。"
        ]
    },
    {
        "title": "整改复核 Agent 是怎么做的？",
        "why": "考察你是否理解闭环，而不只是前半段评分报告。",
        "answer": [
            "整改复核是闭环的关键。传统流程里报告出来后，问题是否整改、整改是否合格，往往靠人工追踪。系统里每个问题会形成整改任务，项目人员提交整改说明和照片，AI 先做初步核验，主管再处理高风险或不确定项。",
            "设计上它不是只看文字，也要结合图片。图片侧可以用视觉模型判断整改前后是否有明显改善，文字侧判断整改说明是否回应了原问题。最终输出通过、驳回或待人工审核，并给出原因和置信度。",
            "产品价值在于把“发现问题”延伸到“问题闭环”。简历里闭环率从不足 40% 提升到 100%，核心不是 AI 写了报告，而是任务、责任人、整改材料、审核结果都结构化沉淀了。"
        ],
        "followups": [
            "如果问风险：图片角度、光线、遮挡会影响判断，所以复核 Agent 只能做辅助审核，高风险项必须人工确认。",
            "如果问模型：可用 Qwen-VL 做图片理解，Qwen-plus/规则做文字与状态判断。"
        ]
    },
    {
        "title": "综合分析 Agent 的价值是什么？",
        "why": "考察你是否能从单次检查走向管理决策。",
        "answer": [
            "综合分析 Agent 解决的是管理层问题，不是现场检查问题。单次报告告诉你某个项目这次哪里不好；综合分析要回答跨项目、跨时间、跨区域哪里更差、问题是否复发、整改是否有效、哪些模块需要专项治理。",
            "系统里综合分析支持三种模式：跨项目对比、跨时间对比、全项目概览。前端用 mode 参数选择 cross_project、cross_time、all_projects，后端做参数校验后进入分析工作流。",
            "LangGraph 里先收集报告数据，再并行分析得分、问题和模块维度，然后用 DeepSeek 生成洞察，最后渲染图表、组装报告、导出文件并保存结果。它的价值是把检查数据从“事后报告”升级成“管理看板和经营分析”。"
        ],
        "followups": [
            "如果问为什么还要 DeepSeek：因为跨报告比较需要长文本归纳和管理建议，不只是简单 SQL 统计。",
            "如果问哪些数据不用 LLM：排名、得分差值、问题数量、趋势图表都应该用规则和 SQL 计算。"
        ]
    },
    {
        "title": "你怎么做模型调用的稳定性保障？",
        "why": "考察生产化意识。",
        "answer": [
            "我从连接、调用、解析、失败兜底四层做稳定性。连接层使用共享 httpx 客户端和连接池，避免频繁建立连接。调用层设置超时、重试、指数退避，并对限流单独处理。解析层强制 JSON 输出，并用 JSON array 提取函数处理模型可能包裹代码块的问题。",
            "失败兜底方面，Qwen-plus 失败后可以降级到 qwen-turbo；如果模型仍失败，系统会基于规则做回退评分，按严重/一般/轻微问题数量扣分。DeepSeek 报告生成也有重试和降级逻辑。",
            "产品上还要把失败显性化，而不是静默吞掉。比如评分异常要在 scoring_basis 中标记，低置信度要进入人工复核，调用耗时、Token、成功率、错误类型进入 LLM 监控看板。"
        ],
        "followups": [
            "如果问为什么不用直接失败提示用户：现场检查场景不能因为模型短暂不可用就中断流程，必须能保存数据并稍后重试。",
            "如果问缓存：相同输入的重跑可以加入缓存，节约成本并提高速度。"
        ]
    },
    {
        "title": "成本控制怎么做？",
        "why": "考察你是否具备商业化和运营视角。",
        "answer": [
            "成本控制首先来自任务分层。数据采集不调用 LLM；合格项不进入 LLM；评分只把有问题项按模块批量发送；报告和综合分析属于低频场景，用更强模型。",
            "第二是 Prompt 控制。评分 Prompt 只保留关键字段，检查标准和评分规则做截断，单批超过上限会分批处理，避免一次调用过长超时。报告分析也只传 item_id、item_name、score、scoring_basis、improvement_suggestion 等摘要字段，不把全部原始数据塞进去。",
            "第三是监控和优化。系统有模型调用与成本监控，看模型类型、Token 消耗、处理耗时、缓存命中率、错误率。后续可以根据监控结果做缓存、模型降级、批量策略调整和 Prompt 压缩。"
        ],
        "followups": [
            "如果问为什么还用 DeepSeek：低频高价值任务用强模型，整体成本仍可控。",
            "如果问 ROI：报告撰写和整改审核节省的人力时间，是成本收益的核心。"
        ]
    },
    {
        "title": "你怎么评估这个 AI 系统做得好不好？",
        "why": "考察你是否会建立 AI 产品指标体系。",
        "answer": [
            "我会分业务指标、AI 效果指标、工程稳定性指标和成本指标四类。业务指标包括检查完成时间、报告生成时间、整改闭环率、整改审核周期、用户满意度。这个项目简历里最关键的业务结果是闭环率从不足 40% 到 100%，全流程从 3-5 天压缩到半天。",
            "AI 效果指标包括评分一致率、人工修改率、严重问题召回率、Precision、Recall、F1、置信度分布、同类问题评分偏差。评分 Agent 的核心不是平均分多高，而是同类问题口径是否一致、高风险问题是否不漏判。",
            "工程稳定性指标包括 API 成功率、JSON 解析成功率、模型超时率、重试成功率、降级次数、前端任务进度准确性。成本指标包括 Token 消耗、单任务模型成本、缓存命中率和平均处理耗时。",
            "这些指标共同决定系统是否可上线。如果只看生成效果，不看人工修改率和严重问题漏判，就不适合企业场景。"
        ],
        "followups": [
            "如果问 F1 怎么用：可把人工复核结果作为标签，统计 AI 对严重/一般/轻微问题识别和整改通过判断的准确性。",
            "如果问没有大量标注怎么办：先用人工复核数据滚动积累 golden dataset。"
        ]
    },
    {
        "title": "如果面试官说你这个不算 Multi-Agent，你怎么回应？",
        "why": "考察你是否能承认边界，同时守住设计合理性。",
        "answer": [
            "我会接受这个挑战，并澄清定义：如果把 Multi-Agent 理解成多个 Agent 自主对话、互相协商，那我的系统不是这种开放式 Multi-Agent；我的系统更准确地说是“多专职 Agent + 工作流编排”。",
            "但从产品和工程角度，它确实有多个独立 AI 能力单元：评分、报告、整改复核、综合分析。它们有不同模型、不同输入输出、不同评估指标、不同失败兜底，并通过数据库和 Orchestrator 串联。",
            "我不追求形式上的 Agent 聊天，而是追求业务流程的可靠落地。品质检查是强合规、强审计场景，所以确定性工作流比开放式协作更合适。这个回答反而能体现我不是盲目追热点，而是按业务选择架构。"
        ],
        "followups": [
            "如果问你会不会 CrewAI：回答“会根据场景选，如果是开放研究或多角色讨论可以用；这个项目不适合”。",
            "如果问简历是否应改：可以把 Multi-Agent 协同具体写成 LangGraph workflow-based agents。"
        ]
    },
    {
        "title": "你在这个项目中最核心的产品决策是什么？",
        "why": "考察你是否能提炼自己的产品判断。",
        "answer": [
            "最核心的产品决策是把 AI 放在“评分、报告、复核、分析”这些高认知成本环节，而不是把所有环节都 AI 化。数据采集、任务分配、权限管理、总分计算都应该规则化，因为这些环节需要确定性。",
            "第二个关键决策是检查完成后统一评分，而不是现场每填一项就调用 AI。这样现场流程不中断，支持离线检查，AI 可以拿到更完整上下文，同时 API 调用次数更少。",
            "第三个关键决策是 AI 结果必须可复核。评分依据、整改建议、置信度、人工修改原因、调用日志都要留下来。这个项目面向的是物业品质管理，不是娱乐生成，所以可追溯比炫技更重要。"
        ],
        "followups": [
            "如果问产品取舍：回答“牺牲了一点实时性，换来现场体验、成本和评分稳定性”。",
            "如果问你的贡献：强调你负责了流程抽象、对象建模、模型分工和评估闭环。"
        ]
    },
    {
        "title": "为什么检查完成后统一评分，而不是边检查边评分？",
        "why": "考察你是否能解释交互体验和系统成本的取舍。",
        "answer": [
            "我选择检查完成后统一评分，主要是因为现场检查场景不能被 AI 调用打断。检查员在现场走动、拍照、记录问题，如果每一项都等模型返回，会严重影响作业效率，也会受网络影响。",
            "统一评分还有三个好处。第一，模型有完整上下文，可以综合判断同一模块下的问题严重程度。第二，按模块批量调用，成本和耗时都更可控。第三，评分完成后主管可以集中复核，不需要在检查过程中频繁处理 AI 结果。",
            "当然它的代价是不能实时给检查员评分反馈。但这个项目的核心目标不是现场即时打分，而是标准化检查、统一评分和自动报告，所以这个取舍是合理的。"
        ],
        "followups": [
            "如果问未来优化：可以对高风险问题做实时提醒，普通评分仍在检查完成后批量处理。",
            "如果问离线：Agent 1 数据采集可离线，Agent 2/3 需要联网调用 LLM。"
        ]
    },
    {
        "title": "数据模型里最重要的对象有哪些？",
        "why": "考察你是否做过对象建模，而不只是写页面需求。",
        "answer": [
            "核心对象可以按流程理解。InspectionTask 是一次完整检查的父任务，包含项目、检查日期、标准类型和总分。InspectionRecord 是某个模块的检查记录，关联任务、模块和检查人员。Issue 是具体问题，关联 record_id、module_name、item_id、描述、位置、严重程度和照片。",
            "ScoringResult 是评分结果，记录 item_id、item_name、score、weight、weighted_score、scoring_basis、improvement_suggestion。Report 是最终报告，保存结构化报告和导出文件。整改模块还会有 Rectification 类对象，用于跟踪提交、AI 审核、人工审核和状态流转。",
            "这些对象的设计核心是可追溯：报告结论能追溯到评分结果，评分结果能追溯到问题和检查项，问题能追溯到照片和现场人员。"
        ],
        "followups": [
            "如果问为什么不用纯 JSON 文件：数据库支持查询、权限、审计、状态恢复和跨项目分析。",
            "如果问模板数据：检查项标准从 Excel 模板解析，并落到统一字段 item_id、item_name、check_standard、check_method、scoring_rule、weight。"
        ]
    },
    {
        "title": "你如何处理检查表模板变化？",
        "why": "考察系统扩展性。",
        "answer": [
            "模板变化是这个项目的核心风险之一，因为物业检查标准可能会调整。我的设计是按列名映射解析 Excel，而不是写死列序号。系统关心的字段包括序号、检查点、检查标准与依据、抽样标准与检查方法、评价方法、权重。",
            "解析时要处理合并单元格、空行、权重格式转换，并按 Sheet 名称识别模块。这样只要模板仍保留标准字段名，列顺序变化不会影响系统。",
            "产品上还需要做模板版本管理。每个 InspectionTask 记录 standard_type 或模板版本，保证历史任务按当时标准评分，不能因为后续模板改了影响历史报告。"
        ],
        "followups": [
            "如果问更复杂变化：新增模块或字段需要模板校验和映射配置页。",
            "如果问为什么模板不直接写进代码：业务标准变更频繁，写死代码维护成本高。"
        ]
    },
    {
        "title": "你怎么防止模型幻觉？",
        "why": "考察 AI 产品安全边界。",
        "answer": [
            "我主要从输入约束、输出约束和职责边界三方面控制。输入上，评分模型只能看到检查项标准、评分规则和用户记录的问题，不让它自由扩展事实。报告模型读取的是数据库里的评分结果和问题摘要，不让它凭空生成检查数据。",
            "输出上，评分必须返回 JSON，字段固定，分数范围限制在 0-5，后端还会做校验。项目总分、模块得分、权重计算不交给模型，而是由后端公式完成。",
            "职责边界上，模型只做语义判断、原因归纳和建议生成；涉及权限、流程状态、总分计算、文件导出、任务流转都由系统规则控制。这样即使模型表达有偏差，也不会破坏核心业务数据。"
        ],
        "followups": [
            "如果问报告幻觉：关键数字来自系统计算，报告只是引用和解释，不能自由编造分数。",
            "如果问审核：低置信度和严重问题进入人工复核。"
        ]
    },
    {
        "title": "这个项目里你做了哪些工程化能力？",
        "why": "考察简历中“全栈落地”和“生产部署”是否可信。",
        "answer": [
            "前端是 Vue3，移动端用 Vant，PC 管理端用 Element Plus。移动端用于现场检查、拍照、问题记录、整改提交；PC 端用于任务管理、评分复核、报告查看、综合分析和系统管理。",
            "后端是 FastAPI，提供认证、任务、检查、评分、报告、整改、综合分析等 REST API。数据库用 SQLite，适合本地部署和轻量化交付。Agent 编排使用 LangGraph，LLM 调用封装在 core/llm_client.py。",
            "工程化上做了 JWT 权限、文件上传、照片压缩、离线同步、进度缓存、SSE 推送、模型调用日志、成本看板、重试熔断、降级模型、导出 Word/PDF 等能力。对 AI 产品经理来说，这些不是炫技，而是保证产品能在真实业务环境跑起来。"
        ],
        "followups": [
            "如果问你是否亲自写代码：按真实情况说“我能用 Vibe Coding 做原型、联调和部署，复杂生产代码与研发协同”。",
            "如果问 Docker：回答“用于简化部署和环境一致性，但本地 SQLite 和文件存储降低了运维门槛”。"
        ]
    },
    {
        "title": "如果重新做一版，你会怎么优化？",
        "why": "考察复盘能力和产品演进判断。",
        "answer": [
            "第一，我会加强 Prompt 和模型策略的版本管理。当前已有 Prompt 管理思路，但如果进入规模化使用，需要支持版本发布、A/B 测试、灰度、回滚和效果对比。",
            "第二，我会建立更标准的 golden dataset。把人工复核、严重问题、边缘案例沉淀为评测集，每次修改 Prompt 或模型前后都跑一致率、严重问题召回率、人工修改率。",
            "第三，我会优化缓存和批处理。相同检查数据重复触发评分时，可以用缓存减少模型成本；整改复核如果数量很多，可以做批量审核。",
            "第四，我会把综合分析从单纯报告输出升级为管理策略建议，比如自动识别高频复发问题、项目排名异常、整改超期风险，并形成专项治理任务。"
        ],
        "followups": [
            "如果问最大遗憾：可以说“早期更关注功能闭环，后续应更早建设评测集和 Prompt 版本治理”。",
            "如果问商业化：可扩展到更多物业项目、第三方检查标准和集团级质量看板。"
        ]
    },
    {
        "title": "面试官质疑你的数据结果，比如闭环率 100%，你怎么解释？",
        "why": "考察你是否能守住量化成果的可信度。",
        "answer": [
            "我不会把 100% 闭环率解释成“所有问题都完美解决”，而是解释成“所有检查问题都进入了系统化跟踪闭环，有责任人、有状态、有整改材料、有审核结果，不再像过去那样线下流转后丢失”。",
            "过去闭环率不足 40%，主要是因为问题记录、照片、整改通知和复核结果分散在线下或聊天工具里，很多问题没有统一状态。系统上线后，每个问题都沉淀为整改任务，状态从待整改、已提交、AI 核查中、AI 通过/驳回、待人工审核、已通过等阶段流转，所以流程闭环率可以提升到 100%。",
            "如果谈业务质量，还要看整改一次通过率、逾期率、复发率、严重问题处理时长等指标。闭环率只是流程数字化的结果，不代表质量治理已经完全完成。"
        ],
        "followups": [
            "如果问报告 3 分钟怎么来的：回答“AI 生成和导出流程自动化后，人工撰写从 2-3 天缩短到分钟级，具体耗时受模块数量和模型响应影响”。",
            "如果问半天周期：解释为从检查、评分、报告到整改任务生成的主流程周期。"
        ]
    },
    {
        "title": "你作为 AI 产品经理，在这个项目里和算法/研发怎么协作？",
        "why": "考察产品经理边界和协作方式。",
        "answer": [
            "我会把协作分成四个层次。第一是业务抽象，把线下检查 SOP、评分规则、角色权限和整改流程整理成 PRD、流程图、对象模型和接口输入输出。第二是 AI 策略，把哪些环节用规则、哪些环节用 LLM、用什么模型、什么 Prompt、什么兜底机制定义清楚。",
            "第三是研发落地，和研发确认 API、数据库表、状态流转、前端交互、导出格式和部署方案。第四是效果迭代，和业务一起做样本复核，和算法/工程一起看错误案例、Prompt 版本、模型成本和指标变化。",
            "我不会把 AI 产品经理理解成只写一句“接入大模型”，而是要把模型能力产品化成可用、可评估、可复核、可持续优化的业务流程。"
        ],
        "followups": [
            "如果问你不会训练模型怎么办：回答“这个项目重点不是训练基础模型，而是模型调用、上下文构造、规则治理和评估闭环”。",
            "如果问你和研发边界：产品负责定义问题、流程、指标和验收标准，研发负责工程实现，双方共同排查效果问题。"
        ]
    },
]

add_heading(doc, "三、详细面试问答", level=1)
for i, q in enumerate(questions, 1):
    add_question(doc, i, q["title"], q["why"], q["answer"], q.get("followups"), q.get("traps"))

add_heading(doc, "四、高频追问速记表", level=1)
quick_rows = [
    ("为什么 Multi-Agent", "不是为了炫技，而是评分、报告、复核、分析的输入输出和风险不同，需要单职责拆分。", "单职责、故障隔离、上下文控制"),
    ("为什么 LangGraph", "这是确定性业务流程，需要状态机、条件边、fan-out/fan-in 和进度追踪。", "StateGraph、可观测、人工门控"),
    ("为什么 DeepSeek", "报告和综合分析是低频长文本综合任务，更看重推理和表达质量。", "长文本、综合归纳、报告质量"),
    ("为什么 Qwen-plus", "评分是高频中文结构化任务，更看重稳定 JSON、规则理解和成本。", "结构化评分、低成本、中文规则"),
    ("意图识别", "主流程不用 LLM 猜，靠角色、入口、状态和 mode 确定性路由；LLM 只识别异常问题描述。", "规则路由、状态机、置信度"),
    ("评分不准", "用结构化输出、置信度、人工复核、历史案例注入和规则回退治理。", "confidence、人工兜底、纠偏闭环"),
    ("成本控制", "合格项不进 LLM，按模块批量，只把问题项和摘要字段发给模型。", "批量、摘要、缓存"),
    ("闭环率 100%", "指所有问题进入系统化状态流转和复核，不代表所有问题天然一次整改成功。", "状态闭环、责任人、复核结果"),
]
for row in quick_rows:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(5)
    r1 = p.add_run(f"{row[0]}：")
    r1.bold = True
    r1.font.name = "Microsoft YaHei"
    r1._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r1.font.size = Pt(10)
    r2 = p.add_run(f"{row[1]} 关键词：{row[2]}")
    r2.font.name = "Microsoft YaHei"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r2.font.size = Pt(10)

add_heading(doc, "五、面试前背诵版", level=1)
add_para(doc, "如果面试官集中问“为什么这么设计”，你可以用下面这段作为总回答：")
add_para(doc, "这个项目我不是按技术热点来设计，而是按物业品质检查的业务流程来拆。数据采集是确定性表单和照片，所以不用 LLM；评分是高频结构化任务，所以用 Qwen-plus，并通过评分规则、Few-shot、历史案例和置信度控制一致性；报告和综合分析是长文本总结和管理洞察，所以用 DeepSeek；整个流程用 LangGraph 做状态编排，用数据库做 Agent 间通信，保证可追溯、可重试、可复核。意图识别上，我没有把它做成开放式聊天机器人，而是用角色、入口、任务状态和模式参数做确定性路由，只有问题描述异常识别这种语义判断交给模型。最终目标不是展示 AI 技术，而是把检查、评分、报告、整改、复核做成一个可闭环、可评估、可持续优化的 AI 产品。")

doc.save(OUTPUT)
print(OUTPUT)
