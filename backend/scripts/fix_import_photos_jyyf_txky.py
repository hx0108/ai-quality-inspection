"""修复金域悦府和田心康苑的问题点照片：从Excel提取嵌入图片并导入数据库"""
import sqlite3, sys, os, json, uuid, io
from datetime import datetime

import openpyxl
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

BASE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
DB_PATH = os.path.join(BASE_DIR, 'data', 'inspection.db')
PHOTOS_DIR = os.path.join(BASE_DIR, 'data', 'photos')
TEMPLATE_PATH = os.path.join(BASE_DIR, 'data', 'templates', '内审检查表V3.0.xlsx')

# 当前数据库中的 task_id
TASK_MAP = {
    '金域悦府': 'Q-20260601-963a3c44',
    '田心康苑': 'Q-20260601-b401e3f3',
}
EXCEL_MAP = {
    '金域悦府': r'c:\Users\ASUS\Desktop\AI Agent\F55金域悦府（蝶城版）品质检查报告.xlsx',
    '田心康苑': r'c:\Users\ASUS\Desktop\AI Agent\F55田心康苑（蝶城版）品质检查报告.xlsx',
}
PREFIX_MAP = {
    '客户服务': '客户', '安全管理': '安全', 'EHS及风险管理': 'EH',
    '环境管理': '环境', '机电运维': '机电', '设施维护': '设施',
    '综合管理': '综合', '财务管理': '财务',
}

# 模板行号到 item_id 的映射（与 import_jyyf_txky.py 保持一致）
# import_jyyf_txky 用 row_idx 从 1 开始编号
def load_template_item_ids(module_name):
    """返回 [(1-based-row-in-excel-data, item_id), ...]
    注意：import_jyyf_txky 中 enumerate(start=1) 从第2行开始读，
    所以第1个数据的 row_idx=1, item_id = prefix-001"""
    wb = openpyxl.load_workbook(TEMPLATE_PATH, read_only=True)
    ws = wb[module_name]
    prefix = PREFIX_MAP[module_name]
    items = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if not row[0]:
            continue
        item_id = f"{prefix}-{row_idx:03d}"
        items.append({'data_row_idx': row_idx, 'item_id': item_id})
    wb.close()
    return items


def import_photos_for_project(project_name):
    task_id = TASK_MAP[project_name]
    excel_path = EXCEL_MAP[project_name]

    print(f'\n{"="*60}')
    print(f'{project_name} | task={task_id}')
    print(f'Excel: {os.path.basename(excel_path)}')
    print(f'{"="*60}')

    if not os.path.exists(excel_path):
        print(f'ERROR: {excel_path} not found')
        return 0

    db = sqlite3.connect(DB_PATH)
    c = db.cursor()

    wb = openpyxl.load_workbook(excel_path)
    total_photos = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        images = ws._images if hasattr(ws, '_images') else []
        if not images:
            continue

        # Get record_id for this module
        c.execute(
            'SELECT record_id FROM inspection_records WHERE task_id=? AND module_name=?',
            (task_id, sheet_name),
        )
        rec = c.fetchone()
        if not rec:
            print(f'  SKIP {sheet_name}: no record in DB')
            continue
        record_id = rec[0]

        # Get issues for this record, indexed by item_id
        c.execute('SELECT issue_id, item_id FROM issues WHERE record_id=?', (record_id,))
        issues_by_item = {}
        for issue_id, item_id in c.fetchall():
            issues_by_item[item_id] = issue_id

        # Load template to map image row positions to item_ids
        # Excel 图片 anchor.row 是 0-indexed，+1 后是 Excel 行号（1-indexed）
        # Excel 数据从第2行开始（第1行是表头），所以 data_row_idx = excel_row - 1
        template_items = load_template_item_ids(sheet_name)
        row_to_item = {}
        for t in template_items:
            # data_row_idx 1 → Excel row 2, data_row_idx 2 → Excel row 3, ...
            excel_row = t['data_row_idx'] + 1
            row_to_item[excel_row] = t['item_id']

        # Skip if already has photos
        c.execute(
            'SELECT COUNT(*) FROM photos WHERE issue_id IN '
            '(SELECT issue_id FROM issues WHERE record_id=?)',
            (record_id,),
        )
        existing = c.fetchone()[0]
        if existing > 0:
            print(f'  SKIP {sheet_name}: already has {existing} photos')
            continue

        # Create photos directory
        photo_dir = os.path.join(PHOTOS_DIR, record_id)
        os.makedirs(photo_dir, exist_ok=True)

        module_photos = 0
        skipped_no_issue = 0
        for img_idx, img in enumerate(images):
            try:
                anchor = img.anchor
                if not hasattr(anchor, '_from'):
                    continue
                excel_row = anchor._from.row + 1  # 0-based to 1-based Excel row
                item_id = row_to_item.get(excel_row)

                if not item_id or item_id not in issues_by_item:
                    skipped_no_issue += 1
                    continue

                issue_id = issues_by_item[item_id]

                # Extract image data
                img_data = img._data()
                img_bytes = img_data.read() if hasattr(img_data, 'read') else img_data

                # Save image
                filename = f'issue_{item_id}_{module_photos}.jpg'
                filepath = os.path.join(photo_dir, filename)

                try:
                    pil_img = Image.open(io.BytesIO(img_bytes))
                    pil_img.save(filepath, 'JPEG', quality=85)
                except Exception:
                    with open(filepath, 'wb') as f:
                        f.write(img_bytes)

                # Create Photo record
                today = datetime.now().strftime('%Y%m%d')
                photo_id = f'PHO-{today}-{uuid.uuid4().hex[:8]}'
                now_str = datetime.now().isoformat()
                c.execute(
                    'INSERT INTO photos (photo_id, issue_id, file_path, file_name, photo_type, upload_time) '
                    'VALUES (?, ?, ?, ?, ?, ?)',
                    (photo_id, issue_id, filepath, filename, '问题照片', now_str),
                )
                module_photos += 1
                total_photos += 1
            except Exception as e:
                print(f'  WARN: image error in {sheet_name} img[{img_idx}]: {e}')

        status = f'{module_photos} photos imported'
        if skipped_no_issue:
            status += f', {skipped_no_issue} skipped (no matching issue)'
        print(f'  {sheet_name}: {len(images)} images in Excel -> {status}')

    db.commit()
    db.close()
    wb.close()
    print(f'Total photos imported: {total_photos}')
    return total_photos


if __name__ == '__main__':
    grand_total = 0
    for name in ['金域悦府', '田心康苑']:
        grand_total += import_photos_for_project(name)

    print(f'\n{"="*60}')
    print(f'ALL DONE. Grand total: {grand_total} photos')
    print(f'{"="*60}')
