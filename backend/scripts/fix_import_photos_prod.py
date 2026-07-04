"""修复金域悦府和田心康苑的问题点照片（生产环境版）
用法:
  1. 将此脚本和两个Excel文件上传到服务器
  2. docker cp 到容器内执行:
     docker cp fix_import_photos_prod.py <容器名>:/tmp/
     docker cp "F55金域悦府（蝶城版）品质检查报告.xlsx" <容器名>:/tmp/F55金域悦府.xlsx
     docker cp "F55田心康苑（蝶城版）品质检查报告.xlsx" <容器名>:/tmp/F55田心康苑.xlsx
  3. docker exec <容器名> python /tmp/fix_import_photos_prod.py
"""
import sqlite3, sys, os, uuid, io
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

try:
    import openpyxl
    from PIL import Image
except ImportError:
    print('ERROR: openpyxl and Pillow are required')
    print('Run: pip install openpyxl Pillow')
    sys.exit(1)

# ============ 生产环境路径 ============
DB_PATH = os.environ.get('DB_PATH', '/app/data/inspection.db')
PHOTOS_DIR = os.environ.get('PHOTOS_DIR', '/app/data/photos')
TEMPLATE_PATH = os.environ.get('TEMPLATE_PATH', '/app/data/templates/内审检查表V3.0.xlsx')

# 生产环境 task_id
TASK_MAP = {
    '金域悦府': 'Q-20260318-c2c1edef',
    '田心康苑': 'Q-20260318-41b967ca',
}
EXCEL_MAP = {
    '金域悦府': '/tmp/F55金域悦府.xlsx',
    '田心康苑': '/tmp/F55田心康苑.xlsx',
}
PREFIX_MAP = {
    '客户服务': '客户', '安全管理': '安全', 'EHS及风险管理': 'EH',
    '环境管理': '环境', '机电运维': '机电', '设施维护': '设施',
    '综合管理': '综合', '财务管理': '财务',
}


def load_template_item_ids(module_name):
    """加载模板，返回 [{data_row_idx, item_id}, ...]"""
    wb = openpyxl.load_workbook(TEMPLATE_PATH, read_only=True)
    ws = wb[module_name]
    prefix = PREFIX_MAP[module_name]
    items = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if not row[0]:
            continue
        item_id = f"{prefix}-{row_idx:03d}"
        items.append({'data_row_idx': row_idx, 'item_id': item_id})
    wb.close()
    return items


def import_photos_for_project(project_name):
    task_id = TASK_MAP[project_name]
    excel_path = EXCEL_MAP[project_name]

    print(f'\n{"="*60}')
    print(f'{project_name} | task={task_id}')
    print(f'{"="*60}')

    if not os.path.exists(excel_path):
        print(f'ERROR: {excel_path} not found')
        return 0

    db = sqlite3.connect(DB_PATH)
    c = db.cursor()

    wb = openpyxl.load_workbook(excel_path)
    total_photos = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        images = ws._images if hasattr(ws, '_images') else []
        if not images:
            continue

        # Get record_id
        c.execute(
            'SELECT record_id FROM inspection_records WHERE task_id=? AND module_name=?',
            (task_id, sheet_name),
        )
        rec = c.fetchone()
        if not rec:
            print(f'  SKIP {sheet_name}: no record')
            continue
        record_id = rec[0]

        # Get issues by item_id
        c.execute('SELECT issue_id, item_id FROM issues WHERE record_id=?', (record_id,))
        issues_by_item = {}
        for issue_id, item_id in c.fetchall():
            issues_by_item[item_id] = issue_id

        # Map template row -> item_id
        template_items = load_template_item_ids(sheet_name)
        row_to_item = {}
        for t in template_items:
            excel_row = t['data_row_idx'] + 1
            row_to_item[excel_row] = t['item_id']

        # Skip if already has photos
        c.execute(
            'SELECT COUNT(*) FROM photos WHERE issue_id IN '
            '(SELECT issue_id FROM issues WHERE record_id=?)',
            (record_id,),
        )
        existing = c.fetchone()[0]
        if existing > 0:
            print(f'  SKIP {sheet_name}: already has {existing} photos')
            continue

        # Create photos directory
        photo_dir = os.path.join(PHOTOS_DIR, record_id)
        os.makedirs(photo_dir, exist_ok=True)

        module_photos = 0
        skipped_no_issue = 0
        for img_idx, img in enumerate(images):
            try:
                anchor = img.anchor
                if not hasattr(anchor, '_from'):
                    continue
                excel_row = anchor._from.row + 1
                item_id = row_to_item.get(excel_row)

                if not item_id or item_id not in issues_by_item:
                    skipped_no_issue += 1
                    continue

                issue_id = issues_by_item[item_id]

                # Extract image
                img_data = img._data()
                img_bytes = img_data.read() if hasattr(img_data, 'read') else img_data

                filename = f'issue_{item_id}_{module_photos}.jpg'
                filepath = os.path.join(photo_dir, filename)

                try:
                    pil_img = Image.open(io.BytesIO(img_bytes))
                    pil_img.save(filepath, 'JPEG', quality=85)
                except Exception:
                    with open(filepath, 'wb') as f:
                        f.write(img_bytes)

                # Create Photo record
                today = datetime.now().strftime('%Y%m%d')
                photo_id = f'PHO-{today}-{uuid.uuid4().hex[:8]}'
                now_str = datetime.now().isoformat()
                c.execute(
                    'INSERT INTO photos (photo_id, issue_id, file_path, file_name, photo_type, upload_time) '
                    'VALUES (?, ?, ?, ?, ?, ?)',
                    (photo_id, issue_id, filepath, filename, '问题照片', now_str),
                )
                module_photos += 1
                total_photos += 1
            except Exception as e:
                print(f'  WARN: {sheet_name} img[{img_idx}]: {e}')

        status = f'{module_photos} photos imported'
        if skipped_no_issue:
            status += f', {skipped_no_issue} skipped'
        print(f'  {sheet_name}: {len(images)} images -> {status}')

    db.commit()
    db.close()
    wb.close()
    print(f'Total: {total_photos}')
    return total_photos


if __name__ == '__main__':
    # Pre-check
    print(f'DB: {DB_PATH} (exists={os.path.exists(DB_PATH)})')
    print(f'Photos dir: {PHOTOS_DIR} (exists={os.path.exists(PHOTOS_DIR)})')
    print(f'Template: {TEMPLATE_PATH} (exists={os.path.exists(TEMPLATE_PATH)})')

    grand_total = 0
    for name in ['金域悦府', '田心康苑']:
        grand_total += import_photos_for_project(name)

    print(f'\n{"="*60}')
    print(f'DONE. Total: {grand_total} photos imported')

    # Restart hint
    if grand_total > 0:
        print('\nIMPORTANT: Restart the backend container to clear in-memory cache:')
        print('  docker restart <container_name>')
    print(f'{"="*60}')
