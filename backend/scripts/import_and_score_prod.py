"""为金域悦府和田心康苑创建检查任务并导入Excel数据 - 生产容器内执行"""
import sys, os, json, uuid
from datetime import datetime
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import sqlite3
import openpyxl

DB_PATH = "/app/data/inspection.db"
TEMPLATE_PATH = "/app/data/templates/内审检查表V3.0.xlsx"
PROJECT_MAP = {"金域悦府": 3, "田心康苑": 2}
MODULE_NAMES = ["客户服务", "安全管理", "EHS及风险管理", "环境管理", "机电运维", "设施维护", "综合管理", "财务管理"]
PREFIX_MAP = {
    "客户服务": "客户", "安全管理": "安全", "EHS及风险管理": "EH",
    "环境管理": "环境", "机电运维": "机电", "设施维护": "设施",
    "综合管理": "综合", "财务管理": "财务",
}

MODULE_WEIGHTS = {
    "客户服务": 0.15, "安全管理": 0.15, "EHS及风险管理": 0.10,
    "环境管理": 0.15, "机电运维": 0.15, "设施维护": 0.15,
    "综合管理": 0.10, "财务管理": 0.05,
}


def gen_id(prefix):
    today = datetime.now().strftime("%Y%m%d")
    return prefix + "-" + today + "-" + uuid.uuid4().hex[:8]


def load_template(module_name):
    wb = openpyxl.load_workbook(TEMPLATE_PATH, read_only=True)
    ws = wb[module_name]
    prefix = PREFIX_MAP[module_name]
    items = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        check_point = row[0]
        if not check_point:
            continue
        item_id = prefix + "-" + str(row_idx).zfill(3)
        items.append({
            "item_id": item_id,
            "item_name": str(check_point),
            "weight": float(row[7]) if row[7] else 0.01,
        })
    wb.close()
    return items


def get_user_issues(module_name, excel_path):
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    if module_name not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[module_name]
    rows_data = []
    if module_name == "财务管理":
        for row in ws.iter_rows(min_row=2, values_only=True):
            desc = str(row[4]).strip() if len(row) > 4 and row[4] else ""
            rows_data.append({"description": desc})
    else:
        for row in ws.iter_rows(min_row=2, values_only=True):
            desc = str(row[8]).strip() if len(row) > 8 and row[8] else ""
            rows_data.append({"description": desc})
    wb.close()
    return rows_data


def import_project(project_name, excel_path):
    project_id = PROJECT_MAP[project_name]
    print("")
    print("=" * 60)
    print("开始导入: " + project_name + " (project_id=" + str(project_id) + ")")
    print("=" * 60)

    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    task_id = gen_id("Q")
    check_date = datetime.now().strftime("%Y-%m-%d")
    now = datetime.now().isoformat()

    c.execute(
        "INSERT INTO inspection_tasks (task_id, project_id, check_date, status, standard_type, created_by, created_at, updated_at) VALUES (?, ?, ?, 'in_progress', 'diecheng', 20, ?, ?)",
        (task_id, project_id, check_date, now, now))
    conn.commit()
    print("Created task: " + task_id)

    total_issues = 0
    total_items = 0

    for module_name in MODULE_NAMES:
        template_items = load_template(module_name)
        user_data = get_user_issues(module_name, excel_path)
        is_ehs = module_name == "EHS及风险管理"
        match_count = min(len(template_items), len(user_data))

        if len(template_items) != len(user_data):
            print("  WARN " + module_name + ": template=" + str(len(template_items)) + " vs excel=" + str(len(user_data)))

        data_json_items = {}
        for i in range(len(template_items)):
            item_id = template_items[i]["item_id"]
            if is_ehs:
                data_json_items[item_id] = {"status": "skipped", "qualified": False}
            elif i < match_count:
                has_issue = bool(user_data[i]["description"] and user_data[i]["description"] != "无")
                data_json_items[item_id] = {"status": "checked", "qualified": not has_issue}
            else:
                data_json_items[item_id] = {"status": "checked", "qualified": True}

        record_id = gen_id("REC")
        c.execute(
            "INSERT INTO inspection_records (record_id, task_id, project_id, module_name, inspector_id, check_date, status, data_json, created_at, updated_at) VALUES (?, ?, ?, ?, 20, ?, 'completed', ?, ?, ?)",
            (record_id, task_id, project_id, module_name, check_date,
             json.dumps({"items": data_json_items}, ensure_ascii=False), now, now))
        total_items += len(data_json_items)

        module_issues = 0
        if not is_ehs:
            for i in range(match_count):
                desc = user_data[i]["description"]
                if not desc or desc == "无":
                    continue
                item = template_items[i]
                severity = "一般"
                issue_id = gen_id("ISS")
                c.execute(
                    "INSERT INTO issues (issue_id, record_id, module_name, item_id, item_name, description, severity, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                    (issue_id, record_id, module_name, item["item_id"], item["item_name"], desc, severity, now))
                module_issues += 1
                total_issues += 1

        print("  " + module_name + ": " + str(len(data_json_items)) + " items, " + str(module_issues) + " issues")

    conn.commit()
    conn.close()
    print("DONE " + project_name + ": " + str(total_items) + " items, " + str(total_issues) + " issues, task_id=" + task_id)
    return task_id


def score_task(task_id):
    print("")
    print("=" * 60)
    print("开始评分: " + task_id)
    print("=" * 60)

    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    now = datetime.now().isoformat()

    c.execute("SELECT record_id, module_name FROM inspection_records WHERE task_id=? ORDER BY id", (task_id,))
    records = c.fetchall()
    print("共 " + str(len(records)) + " 个模块需要评分")

    for record_id, module_name in records:
        c.execute("SELECT COUNT(*) FROM scoring_results WHERE record_id=?", (record_id,))
        if c.fetchone()[0] > 0:
            print("  SKIP " + module_name + ": already scored")
            continue

        template_items = load_template(module_name)
        is_ehs = module_name == "EHS及风险管理"

        c.execute("SELECT item_id, description, severity FROM issues WHERE record_id=?", (record_id,))
        issues_rows = c.fetchall()
        issue_map = {}
        for item_id, desc, severity in issues_rows:
            issue_map[item_id] = {"desc": desc, "severity": severity}

        results_to_insert = []
        for item in template_items:
            item_id = item["item_id"]
            weight = item["weight"]

            if is_ehs:
                score = 5.0
                scoring_basis = "EHS模块本次检查不参与评分"
                is_skipped = True
                is_fallback = False
                suggestion = ""
            elif item_id in issue_map:
                issue = issue_map[item_id]
                severity = issue["severity"]
                if severity == "严重":
                    deduction = 2.0
                elif severity == "轻微":
                    deduction = 0.5
                else:
                    deduction = 1.0
                score = max(0, round(5.0 - deduction, 1))
                scoring_basis = "发现" + severity + "问题：" + issue["desc"][:80] + "，扣" + str(deduction) + "分，得" + str(score) + "分"
                is_skipped = False
                is_fallback = True
                suggestion = "建议制定整改计划并在规定期限内完成"
            else:
                score = 5.0
                scoring_basis = "检查合格，无问题发现，给予满分5分"
                is_skipped = False
                is_fallback = False
                suggestion = ""

            weighted_score = round(score * weight, 6)
            results_to_insert.append((
                "SCR-" + uuid.uuid4().hex[:12],
                record_id, module_name, item_id, item["item_name"],
                score, weight, weighted_score,
                scoring_basis, suggestion,
                is_skipped, now, is_fallback,
            ))

        c.executemany(
            "INSERT INTO scoring_results (scoring_id, record_id, module_name, item_id, item_name, score, weight, weighted_score, scoring_basis, improvement_suggestion, is_skipped, scored_at, is_fallback) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            results_to_insert)
        conn.commit()
        print("  " + module_name + ": " + str(len(results_to_insert)) + " items scored (" + str(len(issue_map)) + " issues)")

    # Calculate module pct scores and total
    c.execute(
        "SELECT sr.module_name, sr.weighted_score, sr.weight FROM scoring_results sr JOIN inspection_records r ON sr.record_id=r.record_id WHERE r.task_id=?",
        (task_id,))
    all_results = c.fetchall()

    modules_data = {}
    for module_name, ws, w in all_results:
        if module_name not in modules_data:
            modules_data[module_name] = {"raw_score_sum": 0, "weight_sum": 0}
        modules_data[module_name]["raw_score_sum"] += float(ws)
        modules_data[module_name]["weight_sum"] += float(w)

    for module_name, data in modules_data.items():
        if data["weight_sum"] > 0:
            pct = round((data["raw_score_sum"] / (5 * data["weight_sum"])) * 100, 2)
            record_ids = [r[0] for r in records if r[1] == module_name]
            for rid in record_ids:
                c.execute("UPDATE scoring_results SET module_pct_score=? WHERE record_id=? AND module_name=?",
                          (pct, rid, module_name))
            print("  " + module_name + " pct_score: " + str(pct))

    project_total = 0
    for module_name, module_weight in MODULE_WEIGHTS.items():
        data = modules_data.get(module_name)
        if data and data["weight_sum"] > 0:
            module_pct = (data["raw_score_sum"] / (5 * data["weight_sum"])) * 100
            project_total += module_pct * module_weight

    project_total = round(project_total, 2)
    c.execute("UPDATE inspection_tasks SET total_score=?, status='completed', updated_at=? WHERE task_id=?",
              (project_total, now, task_id))
    conn.commit()
    print("  ★ 项目总分: " + str(project_total))

    conn.close()
    return project_total


if __name__ == "__main__":
    task_jyyf = import_project("金域悦府", "/tmp/F55金域悦府.xlsx")
    task_txky = import_project("田心康苑", "/tmp/F55田心康苑.xlsx")

    score_jyyf = score_task(task_jyyf)
    score_txky = score_task(task_txky)

    print("")
    print("=" * 60)
    print("金域悦府: task=" + task_jyyf + " score=" + str(score_jyyf))
    print("田心康苑: task=" + task_txky + " score=" + str(score_txky))
    print("=" * 60)

    # Final verification
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT COUNT(*) FROM inspection_tasks')
    print("Total tasks: " + str(c.fetchone()[0]))
    c.execute('SELECT COUNT(*) FROM users')
    print("Total users: " + str(c.fetchone()[0]))
    conn.close()
