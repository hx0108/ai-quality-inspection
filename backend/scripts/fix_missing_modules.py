"""Fix: re-import 5 missing modules for 金域悦府 and 田心康苑, then score"""
import sqlite3, sys, os, json, uuid
from datetime import datetime
import openpyxl

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

DB_PATH = os.environ.get('DB_PATH', '/root/品质检查项目组/data/inspection.db')
TEMPLATE_PATH = os.environ.get('TEMPLATE_PATH', '/root/品质检查项目组/data/templates/内审检查表V3.0.xlsx')

MODULE_NAMES = ['客户服务', '安全管理', 'EHS及风险管理', '环境管理', '机电运维', '设施维护', '综合管理', '财务管理']
MISSING_MODULES = ['客户服务', 'EHS及风险管理', '设施维护', '综合管理', '财务管理']
PREFIX_MAP = {
    '客户服务': '客户', '安全管理': '安全', 'EHS及风险管理': 'EH',
    '环境管理': '环境', '机电运维': '机电', '设施维护': '设施',
    '综合管理': '综合', '财务管理': '财务',
}
MODULE_WEIGHTS = {
    '客户服务': 0.15, '安全管理': 0.15, 'EHS及风险管理': 0.10,
    '环境管理': 0.15, '机电运维': 0.15, '设施维护': 0.15,
    '综合管理': 0.10, '财务管理': 0.05,
}

TASK_MAP = {
    3: 'Q-20260318-c2c1edef',  # 金域悦府
    2: 'Q-20260318-41b967ca',  # 田心康苑
}
EXCEL_MAP = {
    3: '/tmp/F55金域悦府.xlsx',
    2: '/tmp/F55田心康苑.xlsx',
}


def gen_id(prefix):
    today = '20260318'
    return prefix + '-' + today + '-' + uuid.uuid4().hex[:8]


def load_template(module_name):
    wb = openpyxl.load_workbook(TEMPLATE_PATH, read_only=True)
    ws = wb[module_name]
    prefix = PREFIX_MAP[module_name]
    items = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        check_point = row[0]
        if not check_point:
            continue
        item_id = prefix + '-' + str(row_idx).zfill(3)
        items.append({
            'item_id': item_id,
            'item_name': str(check_point),
            'weight': float(row[7]) if row[7] else 0.01,
        })
    wb.close()
    return items


def get_user_issues(module_name, excel_path):
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    if module_name not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[module_name]
    rows_data = []
    if module_name == '财务管理':
        for row in ws.iter_rows(min_row=2, values_only=True):
            desc = str(row[4]).strip() if len(row) > 4 and row[4] else ''
            rows_data.append({'description': desc})
    else:
        for row in ws.iter_rows(min_row=2, values_only=True):
            desc = str(row[8]).strip() if len(row) > 8 and row[8] else ''
            rows_data.append({'description': desc})
    wb.close()
    return rows_data


def fix_project(project_id, excel_path):
    task_id = TASK_MAP[project_id]
    print('\n=== Fixing project_id=' + str(project_id) + ' task_id=' + task_id + ' ===')

    db = sqlite3.connect(DB_PATH)
    c = db.cursor()
    now_str = datetime.now().isoformat()

    # Delete incomplete records + their issues + scoring for missing modules
    for mod in MISSING_MODULES:
        # Find record
        c.execute('SELECT record_id FROM inspection_records WHERE task_id=? AND module_name=?', (task_id, mod))
        rec = c.fetchone()
        if rec:
            record_id = rec[0]
            c.execute('DELETE FROM scoring_results WHERE record_id=?', (record_id,))
            c.execute('DELETE FROM issues WHERE record_id=?', (record_id,))
            c.execute('DELETE FROM inspection_records WHERE record_id=?', (record_id,))
            print('  Cleaned: ' + mod + ' (record=' + record_id + ')')

    db.commit()

    # Re-import missing modules
    total_issues = 0
    for module_name in MISSING_MODULES:
        template_items = load_template(module_name)
        user_data = get_user_issues(module_name, excel_path)
        is_ehs = module_name == 'EHS及风险管理'
        match_count = min(len(template_items), len(user_data))

        data_json_items = {}
        for i in range(len(template_items)):
            item_id = template_items[i]['item_id']
            if is_ehs:
                data_json_items[item_id] = {'status': 'skipped', 'qualified': False}
            elif i < match_count:
                has_issue = bool(user_data[i]['description'] and user_data[i]['description'] != '无')
                data_json_items[item_id] = {'status': 'checked', 'qualified': not has_issue}
            else:
                data_json_items[item_id] = {'status': 'checked', 'qualified': True}

        record_id = gen_id('REC')
        c.execute(
            'INSERT INTO inspection_records (record_id, task_id, project_id, module_name, inspector_id, check_date, status, data_json, created_at, updated_at) VALUES (?, ?, ?, ?, 20, ?, ?, ?, ?, ?)',
            (record_id, task_id, project_id, module_name, '2026-03-18', 'completed',
             json.dumps({'items': data_json_items}, ensure_ascii=False), now_str, now_str))

        # Create issues
        module_issues = 0
        if not is_ehs:
            for i in range(match_count):
                desc = user_data[i]['description']
                if not desc or desc == '无':
                    continue
                item = template_items[i]
                severity = '一般'
                issue_id = gen_id('ISS')
                c.execute(
                    'INSERT INTO issues (issue_id, record_id, module_name, item_id, item_name, description, severity, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)',
                    (issue_id, record_id, module_name, item['item_id'], item['item_name'], desc, severity, now_str))
                module_issues += 1
                total_issues += 1

        print('  Imported: ' + module_name + ' (' + str(len(data_json_items)) + ' items, ' + str(module_issues) + ' issues)')

    db.commit()

    # Score all missing modules
    for module_name in MISSING_MODULES:
        c.execute('SELECT record_id FROM inspection_records WHERE task_id=? AND module_name=?', (task_id, module_name))
        rec = c.fetchone()
        if not rec:
            print('  SKIP score: ' + module_name + ' no record')
            continue
        record_id = rec[0]
        template_items = load_template(module_name)
        is_ehs = module_name == 'EHS及风险管理'

        c.execute('SELECT item_id, description, severity FROM issues WHERE record_id=?', (record_id,))
        issues_rows = c.fetchall()
        issue_map = {}
        for item_id, desc, severity in issues_rows:
            issue_map[item_id] = {'desc': desc, 'severity': severity}

        results = []
        for item in template_items:
            item_id = item['item_id']
            weight = item['weight']
            if is_ehs:
                score = 5.0
                basis = 'EHS模块本次检查不参与评分'
                skipped = True
                fallback = False
                suggestion = ''
            elif item_id in issue_map:
                issue = issue_map[item_id]
                sev = issue['severity']
                deduction = 2.0 if sev == '严重' else (0.5 if sev == '轻微' else 1.0)
                score = max(0, round(5.0 - deduction, 1))
                basis = '发现' + sev + '问题：' + issue['desc'][:80] + '，扣' + str(deduction) + '分，得' + str(score) + '分'
                skipped = False
                fallback = True
                suggestion = '建议制定整改计划并在规定期限内完成'
            else:
                score = 5.0
                basis = '检查合格，无问题发现，给予满分5分'
                skipped = False
                fallback = False
                suggestion = ''
            weighted_score = round(score * weight, 6)
            results.append((
                'SCR-' + uuid.uuid4().hex[:12],
                record_id, module_name, item_id, item['item_name'],
                score, weight, weighted_score,
                basis, suggestion,
                skipped, now_str, fallback,
            ))

        c.executemany(
            'INSERT INTO scoring_results (scoring_id, record_id, module_name, item_id, item_name, score, weight, weighted_score, scoring_basis, improvement_suggestion, is_skipped, scored_at, is_fallback) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)',
            results)
        print('  Scored: ' + module_name + ' (' + str(len(results)) + ' items)')

    db.commit()
    db.close()
    print('  Total issues imported: ' + str(total_issues))


def recompute_total(task_id):
    db = sqlite3.connect(DB_PATH)
    c = db.cursor()
    c.execute(
        'SELECT sr.module_name, sr.weighted_score, sr.weight FROM scoring_results sr JOIN inspection_records r ON sr.record_id=r.record_id WHERE r.task_id=?',
        (task_id,))
    all_results = c.fetchall()
    modules_data = {}
    for module_name, ws, w in all_results:
        if module_name not in modules_data:
            modules_data[module_name] = {'raw': 0, 'weight': 0}
        modules_data[module_name]['raw'] += float(ws)
        modules_data[module_name]['weight'] += float(w)

    for module_name, data in modules_data.items():
        if data['weight'] > 0:
            pct = round((data['raw'] / (5 * data['weight'])) * 100, 2)
            c.execute('UPDATE scoring_results SET module_pct_score=? WHERE record_id IN (SELECT record_id FROM inspection_records WHERE task_id=? AND module_name=?)', (pct, task_id, module_name))
            print('  ' + module_name + ': pct=' + str(pct))

    project_total = 0
    for module_name, module_weight in MODULE_WEIGHTS.items():
        data = modules_data.get(module_name)
        if data and data['weight'] > 0:
            module_pct = (data['raw'] / (5 * data['weight'])) * 100
            project_total += module_pct * module_weight

    project_total = round(project_total, 2)
    c.execute('UPDATE inspection_tasks SET total_score=? WHERE task_id=?', (project_total, task_id))
    db.commit()
    print('  Total score: ' + str(project_total))
    db.close()


if __name__ == '__main__':
    # Re-upload Excel files to /tmp/ first
    for project_id in [3, 2]:
        excel_path = EXCEL_MAP[project_id]
        if not os.path.exists(excel_path):
            print('ERROR: Excel not found: ' + excel_path)
            sys.exit(1)

    fix_project(3, EXCEL_MAP[3])
    fix_project(2, EXCEL_MAP[2])

    print('\n=== Recomputing totals ===')
    recompute_total(TASK_MAP[3])
    recompute_total(TASK_MAP[2])

    # Final verify
    db = sqlite3.connect(DB_PATH)
    c = db.cursor()
    for tid in list(TASK_MAP.values()):
        c.execute('SELECT COUNT(*) FROM inspection_records WHERE task_id=?', (tid,))
        recs = c.fetchone()[0]
        c.execute('SELECT COUNT(DISTINCT sr.module_name) FROM scoring_results sr JOIN inspection_records r ON sr.record_id=r.record_id WHERE r.task_id=?', (tid,))
        scored = c.fetchone()[0]
        c.execute('SELECT total_score FROM inspection_tasks WHERE task_id=?', (tid,))
        score = c.fetchone()[0]
        print('  ' + tid + ': ' + str(recs) + ' records, ' + str(scored) + ' scored modules, total=' + str(score))
    db.close()
    print('\nDone!')
