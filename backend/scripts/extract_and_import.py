"""通用导入脚本：从已完成的品质检查Excel导入检查数据到系统"""
import sys, os, json, uuid, shutil
from datetime import datetime

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, "/app")

from database import SessionLocal
from models.models import InspectionRecord, Issue, Photo
import openpyxl
from PIL import Image
import io

# ==================== 配置 ====================
TASK_ID = sys.argv[1]  # e.g. "Q-20260318-e198ba12"
EXCEL_PATH = sys.argv[2]  # e.g. "/tmp/import_data/F55金域悦府.xlsx"
TEMPLATE_PATH = "/app/data/templates/内审检查表V3.0.xlsx"
PHOTOS_STORAGE = "/app/data/photos"
MODULE_NAMES = ["客户服务", "安全管理", "EHS及风险管理", "环境管理", "机电运维", "设施维护", "综合管理", "财务管理"]


def gen_id(prefix):
    today = datetime.now().strftime("%Y%m%d")
    return f"{prefix}-{today}-{uuid.uuid4().hex[:8]}"


def load_template_item_ids(module_name):
    wb = openpyxl.load_workbook(TEMPLATE_PATH, read_only=True)
    ws = wb[module_name]
    items = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if not row[0]:
            continue
        item_id = f"{module_name[:2]}-{row_idx:03d}"
        items.append({"item_id": item_id, "item_name": str(row[0])})
    wb.close()
    return items


def get_user_excel_data(module_name, excel_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if module_name not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[module_name]
    rows_data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        # Column I (index 8) = description, but some sheets have fewer columns
        desc = str(row[8]).strip() if len(row) > 8 and row[8] else ""
        score = row[10] if len(row) > 10 and row[10] is not None else None
        rows_data.append({"description": desc, "score": score})
    wb.close()
    return rows_data


def extract_images(excel_path, output_dir):
    """Extract all embedded images from Excel and return mapping"""
    os.makedirs(output_dir, exist_ok=True)
    wb = openpyxl.load_workbook(excel_path)
    image_map = {}
    total = 0

    for name in wb.sheetnames:
        ws = wb[name]
        images = ws._images
        if not images:
            continue
        sheet_dir = os.path.join(output_dir, name)
        os.makedirs(sheet_dir, exist_ok=True)
        sheet_map = {}

        for i, img in enumerate(images):
            try:
                anchor = img.anchor
                if not hasattr(anchor, "_from"):
                    continue
                col_idx = anchor._from.col
                row_idx = anchor._from.row
                excel_row = row_idx + 1

                img_data = img._data()
                img_bytes = img_data.read() if hasattr(img_data, "read") else img_data

                filename = f"row{excel_row}_img{i}.jpg"
                filepath = os.path.join(sheet_dir, filename)

                try:
                    pil_img = Image.open(io.BytesIO(img_bytes))
                    pil_img.save(filepath, "JPEG", quality=85)
                except Exception:
                    with open(filepath, "wb") as f:
                        f.write(img_bytes)

                if str(excel_row) not in sheet_map:
                    sheet_map[str(excel_row)] = []
                sheet_map[str(excel_row)].append(filename)
                total += 1
            except Exception as e:
                print(f"  WARN: image {i} in {name}: {e}")

        image_map[name] = sheet_map

    wb.close()
    # Save map
    map_path = os.path.join(output_dir, "image_map.json")
    with open(map_path, "w", encoding="utf-8") as f:
        json.dump(image_map, f, ensure_ascii=False)

    print(f"Images extracted: {total} from {len(image_map)} sheets")
    return image_map


def main():
    db = SessionLocal()
    try:
        task = db.query(InpectionTask if False else InspectionRecord).first()  # just to test import
        from models.models import InspectionTask
        task = db.query(InspectionTask).filter(InspectionTask.task_id == TASK_ID).first()
        if not task:
            print(f"ERROR: Task {TASK_ID} not found")
            return

        existing = db.query(InspectionRecord).filter(InspectionRecord.task_id == TASK_ID).count()
        if existing > 0:
            print(f"Cleanup: removing {existing} existing records...")
            # Delete photos, issues, records for this task
            for rec in db.query(InspectionRecord).filter(InspectionRecord.task_id == TASK_ID).all():
                for issue in db.query(Issue).filter(Issue.record_id == rec.record_id).all():
                    db.query(Photo).filter(Photo.issue_id == issue.issue_id).delete()
                db.query(Issue).filter(Issue.record_id == rec.record_id).delete()
            db.query(InspectionRecord).filter(InspectionRecord.task_id == TASK_ID).delete()
            db.commit()
            print("  Cleaned up")

        print(f"Task: {TASK_ID}, project={task.project_id}")

        # Extract images
        photos_dir = os.path.join("/tmp", f"photos_{TASK_ID}")
        image_map = extract_images(EXCEL_PATH, photos_dir)

        total_issues = 0
        total_photos = 0
        total_items = 0

        for module_name in MODULE_NAMES:
            template_items = load_template_item_ids(module_name)
            user_data = get_user_excel_data(module_name, EXCEL_PATH)
            is_ehs = module_name == "EHS及风险管理"
            match_count = min(len(template_items), len(user_data))

            if len(template_items) != len(user_data):
                print(f"  {module_name}: template={len(template_items)} vs excel={len(user_data)}")

            # Create record
            record_id = gen_id("REC")
            data_json_items = {}
            for i in range(len(template_items)):
                item_id = template_items[i]["item_id"]
                if is_ehs:
                    data_json_items[item_id] = {"status": "skipped", "qualified": False}
                elif i < match_count:
                    has_issue = bool(user_data[i]["description"] and user_data[i]["description"] != "无")
                    data_json_items[item_id] = {"status": "checked", "qualified": not has_issue}
                else:
                    data_json_items[item_id] = {"status": "checked", "qualified": True}

            record = InspectionRecord(
                record_id=record_id,
                task_id=TASK_ID,
                project_id=task.project_id,
                module_name=module_name,
                inspector_id=20,
                check_date=task.check_date,
                status="completed",
                data_json=json.dumps({"items": data_json_items}, ensure_ascii=False),
            )
            db.add(record)
            db.flush()
            total_items += len(data_json_items)

            # Create issues + photos
            module_image_map = image_map.get(module_name, {})
            module_issues = 0

            if not is_ehs:
                for i in range(match_count):
                    desc = user_data[i]["description"]
                    if not desc or desc == "无":
                        continue

                    item_id = template_items[i]["item_id"]
                    item_name = template_items[i]["item_name"]
                    score = user_data[i]["score"]

                    if score is not None and isinstance(score, (int, float)):
                        severity = "严重" if score == 0 else ("一般" if score <= 2 else "轻微")
                    else:
                        severity = "一般"

                    issue_id = gen_id("ISS")
                    issue = Issue(
                        issue_id=issue_id,
                        record_id=record_id,
                        module_name=module_name,
                        item_id=item_id,
                        item_name=item_name,
                        description=desc,
                        severity=severity,
                    )
                    db.add(issue)
                    db.flush()
                    module_issues += 1
                    total_issues += 1

                    # Photos
                    excel_row = i + 2
                    photo_files = module_image_map.get(str(excel_row), [])
                    for pf in photo_files:
                        src = os.path.join(photos_dir, module_name, pf)
                        if not os.path.exists(src):
                            continue
                        photo_dir = os.path.join(PHOTOS_STORAGE, record_id)
                        os.makedirs(photo_dir, exist_ok=True)
                        dst = os.path.join(photo_dir, pf)
                        shutil.copy2(src, dst)
                        photo = Photo(
                            photo_id=gen_id("PHO"),
                            issue_id=issue_id,
                            file_path=dst,
                            file_name=pf,
                            photo_type="问题照片",
                        )
                        db.add(photo)
                        total_photos += 1

            print(f"  {module_name}: {len(data_json_items)} items, {module_issues} issues")

        db.commit()
        print(f"\nDone! {total_items} items, {total_issues} issues, {total_photos} photos")

    except Exception as e:
        db.rollback()
        print(f"ERROR: {e}")
        import traceback
        traceback.print_exc()
    finally:
        db.close()


if __name__ == "__main__":
    main()
