"""为金域悦府和田心康苑创建检查任务并导入Excel数据"""
import sys, os, json, uuid
from datetime import datetime

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import sqlite3
import openpyxl

DB_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "inspection.db")
TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "templates", "内审检查表V3.0.xlsx")

PROJECT_MAP = {"金域悦府": 3, "田心康苑": 2}
MODULE_NAMES = ["客户服务", "安全管理", "EHS及风险管理", "环境管理", "机电运维", "设施维护", "综合管理", "财务管理"]
PREFIX_MAP = {
    "客户服务": "客户", "安全管理": "安全", "EHS及风险管理": "EH",
    "环境管理": "环境", "机电运维": "机电", "设施维护": "设施",
    "综合管理": "综合", "财务管理": "财务",
}


def gen_id(prefix):
    today = datetime.now().strftime("%Y%m%d")
    return f"{prefix}-{today}-{uuid.uuid4().hex[:8]}"


def load_template(module_name):
    """从模板加载检查项：item_id, item_name(检查点), weight, 评价方法, 抽样标准"""
    wb = openpyxl.load_workbook(TEMPLATE_PATH, read_only=True)
    ws = wb[module_name]
    prefix = PREFIX_MAP[module_name]
    items = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        check_point = row[0]
        if not check_point:
            continue
        item_id = f"{prefix}-{row_idx:03d}"
        items.append({
            "item_id": item_id,
            "item_name": str(check_point),
            "weight": float(row[7]) if row[7] else 0.01,
            "scoring_rule": str(row[3])[:200] if row[3] else "",
            "check_method": str(row[2])[:200] if row[2] else "",
            "check_standard": str(row[1])[:500] if row[1] else "",
        })
    wb.close()
    return items


def get_user_issues(module_name, excel_path):
    """从用户Excel提取问题点"""
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    if module_name not in wb.sheetnames:
        wb.close()
        return []

    ws = wb[module_name]
    rows_data = []

    if module_name == "财务管理":
        # 财务管理: col 0=检查点, col 4=问题描述, col 5=图片
        for row in ws.iter_rows(min_row=2, values_only=True):
            desc = str(row[4]).strip() if len(row) > 4 and row[4] else ""
            rows_data.append({"description": desc})
    else:
        # 其他7个模块: col 8=问题描述, col 9=图片
        for row in ws.iter_rows(min_row=2, values_only=True):
            desc = str(row[8]).strip() if len(row) > 8 and row[8] else ""
            rows_data.append({"description": desc})

    wb.close()
    return rows_data


def import_project(project_name, excel_path):
    project_id = PROJECT_MAP[project_name]
    print(f"\n{'='*60}")
    print(f"开始导入: {project_name} (project_id={project_id})")
    print(f"Excel: {os.path.basename(excel_path)}")
    print(f"{'='*60}")

    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()

    # 1. 创建检查任务
    task_id = gen_id("Q")
    check_date = datetime.now().strftime("%Y-%m-%d")
    now = datetime.now().isoformat()

    c.execute("""
        INSERT INTO inspection_tasks (task_id, project_id, check_date, status, standard_type, created_by, created_at, updated_at)
        VALUES (?, ?, ?, 'in_progress', 'diecheng', 20, ?, ?)
    """, (task_id, project_id, check_date, now, now))
    conn.commit()
    print(f"✓ 创建任务: {task_id}, date={check_date}")

    total_issues = 0
    total_items = 0

    # 2. 逐模块导入
    for module_name in MODULE_NAMES:
        template_items = load_template(module_name)
        user_data = get_user_issues(module_name, excel_path)
        is_ehs = module_name == "EHS及风险管理"

        match_count = min(len(template_items), len(user_data))
        if len(template_items) != len(user_data):
            print(f"  ⚠ {module_name}: template={len(template_items)} vs excel={len(user_data)}")

        # 构建 data_json
        data_json_items = {}
        for i in range(len(template_items)):
            item_id = template_items[i]["item_id"]
            if is_ehs:
                data_json_items[item_id] = {"status": "skipped", "qualified": False}
            elif i < match_count:
                has_issue = bool(user_data[i]["description"] and user_data[i]["description"] != "无")
                data_json_items[item_id] = {"status": "checked", "qualified": not has_issue}
            else:
                data_json_items[item_id] = {"status": "checked", "qualified": True}

        # 创建 InspectionRecord
        record_id = gen_id("REC")
        c.execute("""
            INSERT INTO inspection_records (record_id, task_id, project_id, module_name, inspector_id, check_date, status, data_json, created_at, updated_at)
            VALUES (?, ?, ?, ?, 20, ?, 'completed', ?, ?, ?)
        """, (record_id, task_id, project_id, module_name, check_date,
              json.dumps({"items": data_json_items}, ensure_ascii=False), now, now))

        total_items += len(data_json_items)

        # 创建 Issues
        module_issues = 0
        if not is_ehs:
            for i in range(match_count):
                desc = user_data[i]["description"]
                if not desc or desc == "无":
                    continue

                item = template_items[i]
                # 严重程度推断
                severity = "一般"

                issue_id = gen_id("ISS")
                c.execute("""
                    INSERT INTO issues (issue_id, record_id, module_name, item_id, item_name, description, severity, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """, (issue_id, record_id, module_name, item["item_id"], item["item_name"], desc, severity, now))
                module_issues += 1
                total_issues += 1

        print(f"  ✓ {module_name}: {len(data_json_items)} items, {module_issues} issues")

    conn.commit()
    conn.close()
    print(f"\n✓ {project_name} 导入完成: {total_items} items, {total_issues} issues, task_id={task_id}")
    return task_id


if __name__ == "__main__":
    excel_jyyf = r"c:\Users\ASUS\Desktop\AI Agent\F55金域悦府（蝶城版）品质检查报告.xlsx"
    excel_txky = r"c:\Users\ASUS\Desktop\AI Agent\F55田心康苑（蝶城版）品质检查报告.xlsx"

    task_jyyf = import_project("金域悦府", excel_jyyf)
    task_txky = import_project("田心康苑", excel_txky)

    print(f"\n{'='*60}")
    print(f"RESULT: task_jyyf={task_jyyf}")
    print(f"RESULT: task_txky={task_txky}")
    print(f"{'='*60}")
