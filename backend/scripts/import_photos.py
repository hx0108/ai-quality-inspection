"""Extract images from Excel and import as Photo records for 金域悦府 and 田心康苑"""
import sqlite3, sys, os, json, uuid, io
from datetime import datetime
import openpyxl
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

DB_PATH = os.environ.get('DB_PATH', '/app/data/inspection.db')
PHOTOS_DIR = os.environ.get('PHOTOS_DIR', '/app/data/photos')
TEMPLATE_PATH = os.environ.get('TEMPLATE_PATH', '/app/data/templates/内审检查表V3.0.xlsx')

TASK_MAP = {
    3: 'Q-20260318-c2c1edef',  # 金域悦府
    2: 'Q-20260318-41b967ca',  # 田心康苑
}
EXCEL_MAP = {
    3: '/tmp/F55金域悦府.xlsx',
    2: '/tmp/F55田心康苑.xlsx',
}
PREFIX_MAP = {
    '客户服务': '客户', '安全管理': '安全', 'EHS及风险管理': 'EH',
    '环境管理': '环境', '机电运维': '机电', '设施维护': '设施',
    '综合管理': '综合', '财务管理': '财务',
}

def gen_id(prefix):
    return prefix + '-20260318-' + uuid.uuid4().hex[:8]

def load_template_item_ids(module_name):
    wb = openpyxl.load_workbook(TEMPLATE_PATH, read_only=True)
    ws = wb[module_name]
    prefix = PREFIX_MAP[module_name]
    items = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:
            continue
        item_id = prefix + '-' + str(row_idx - 1).zfill(3)
        items.append({'row': row_idx, 'item_id': item_id})
    wb.close()
    return items

def import_photos_for_project(project_id, excel_path):
    task_id = TASK_MAP[project_id]
    print('\n=== Project ' + str(project_id) + ' task=' + task_id + ' ===')

    db = sqlite3.connect(DB_PATH)
    c = db.cursor()

    wb = openpyxl.load_workbook(excel_path)
    total_photos = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        images = ws._images if hasattr(ws, '_images') else []
        if not images:
            continue

        # Get record_id for this module
        c.execute('SELECT record_id FROM inspection_records WHERE task_id=? AND module_name=?', (task_id, sheet_name))
        rec = c.fetchone()
        if not rec:
            print('  SKIP ' + sheet_name + ': no record')
            continue
        record_id = rec[0]

        # Get issues for this record, indexed by item_id
        c.execute('SELECT issue_id, item_id FROM issues WHERE record_id=?', (record_id,))
        issues_by_item = {}
        for issue_id, item_id in c.fetchall():
            issues_by_item[item_id] = issue_id

        # Load template to map row numbers to item_ids
        template_items = load_template_item_ids(sheet_name)
        row_to_item = {}
        for t in template_items:
            row_to_item[t['row']] = t['item_id']

        # Get existing photos count for this record (to avoid duplicates)
        c.execute('SELECT COUNT(*) FROM photos WHERE issue_id IN (SELECT issue_id FROM issues WHERE record_id=?)', (record_id,))
        existing = c.fetchone()[0]
        if existing > 0:
            print('  SKIP ' + sheet_name + ': already has ' + str(existing) + ' photos')
            continue

        # Create photos directory
        photo_dir = os.path.join(PHOTOS_DIR, record_id)
        os.makedirs(photo_dir, exist_ok=True)

        module_photos = 0
        for img in images:
            try:
                anchor = img.anchor
                if not hasattr(anchor, '_from'):
                    continue
                excel_row = anchor._from.row + 1  # 0-based to 1-based
                item_id = row_to_item.get(excel_row)

                if not item_id or item_id not in issues_by_item:
                    # No matching issue for this image
                    continue

                issue_id = issues_by_item[item_id]

                # Extract image data
                img_data = img._data()
                img_bytes = img_data.read() if hasattr(img_data, 'read') else img_data

                # Save image
                filename = 'issue_' + item_id + '_' + str(module_photos) + '.jpg'
                filepath = os.path.join(photo_dir, filename)

                try:
                    pil_img = Image.open(io.BytesIO(img_bytes))
                    pil_img.save(filepath, 'JPEG', quality=85)
                except Exception:
                    with open(filepath, 'wb') as f:
                        f.write(img_bytes)

                # Create Photo record
                photo_id = gen_id('PHO')
                now_str = datetime.now().isoformat()
                c.execute(
                    'INSERT INTO photos (photo_id, issue_id, file_path, file_name, photo_type, upload_time) VALUES (?, ?, ?, ?, ?, ?)',
                    (photo_id, issue_id, filepath, filename, '问题照片', now_str))
                module_photos += 1
                total_photos += 1
            except Exception as e:
                print('  WARN: image error in ' + sheet_name + ': ' + str(e)[:60])

        if module_photos > 0:
            print('  ' + sheet_name + ': ' + str(module_photos) + ' photos imported')

    db.commit()
    db.close()
    wb.close()
    print('Total photos: ' + str(total_photos))
    return total_photos


if __name__ == '__main__':
    total = 0
    for project_id in [3, 2]:
        excel_path = EXCEL_MAP[project_id]
        if not os.path.exists(excel_path):
            print('ERROR: ' + excel_path + ' not found')
            sys.exit(1)
        total += import_photos_for_project(project_id, excel_path)

    print('\nAll done! Total photos imported: ' + str(total))
