"""Excel builders shared by batch scoring exports."""

import io
import os
import tempfile
from collections import defaultdict

import openpyxl
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PilImage, ImageOps

from core import standards
from core.scoring_aggregation import aggregate


STANDARD_LABELS = {
    "diecheng": "蝶城版",
    "feidiecheng": "非蝶城版",
    "lizhi": "砺质版",
}
STATUS_LABELS = {
    "pending": "待开始",
    "in_progress": "进行中",
    "completed": "已完成",
}


def resolve_photo_path(file_path):
    """Resolve absolute and backend-relative photo paths without changing data."""
    if not file_path:
        return None
    backend_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    candidates = [file_path]
    if not os.path.isabs(file_path):
        candidates.extend([
            os.path.join(backend_root, file_path),
            os.path.join(os.getcwd(), file_path),
        ])
    for candidate in candidates:
        normalized = os.path.abspath(candidate)
        if os.path.isfile(normalized):
            return normalized
    return None


def _item_max_score(standard_type, module_name, item_id, template_cache):
    if standard_type != "lizhi":
        return 5.0
    key = (standard_type, module_name)
    if key not in template_cache:
        from api.tasks import load_template_items
        template_cache[key] = {
            str(item.get("item_id")): float(item.get("max_score") or 0)
            for item in load_template_items(module_name, standard_type)
        }
    return float(
        template_cache[key].get(str(item_id))
        or standards.get_module_cfg(standard_type, module_name).get("max_score")
        or 0
    )


def _apply_header(cell, header_font, header_fill, border):
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border


def _write_row(ws, row_index, values, cell_font, border, center_columns=()):
    for column_index, value in enumerate(values, 1):
        cell = ws.cell(row=row_index, column=column_index, value=value)
        cell.font = cell_font
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center" if column_index in center_columns else "left",
            vertical="center",
            wrap_text=True,
        )


def build_batch_scoring_workbook(tasks, records, results, issues, photos):
    """Build a three-sheet scoring workbook and return its temporary path."""
    task_by_id = {task.task_id: task for task in tasks}
    record_by_id = {record.record_id: record for record in records}
    results_by_task = defaultdict(list)
    issues_by_item = defaultdict(list)
    photos_by_item = defaultdict(list)

    for result in results:
        record = record_by_id.get(result.record_id)
        if record and record.task_id in task_by_id:
            results_by_task[record.task_id].append(result)

    issue_by_id = {}
    for issue in issues:
        key = (issue.record_id, issue.module_name, issue.item_id)
        issues_by_item[key].append(issue)
        issue_by_id[issue.issue_id] = issue

    for photo in photos:
        issue = issue_by_id.get(photo.issue_id)
        if issue:
            key = (issue.record_id, issue.module_name, issue.item_id)
            photos_by_item[key].append(photo)

    max_photo_count = max([len(value) for value in photos_by_item.values()] or [0])
    max_photo_count = max(1, max_photo_count)

    wb = openpyxl.Workbook()
    image_streams = []
    ws_tasks = wb.active
    ws_tasks.title = "任务汇总"
    ws_modules = wb.create_sheet("模块汇总")
    ws_details = wb.create_sheet("检查项明细")

    header_font = Font(name="微软雅黑", bold=True, size=10, color="FFFFFF")
    header_fill = PatternFill(start_color="1A1F3C", end_color="1A1F3C", fill_type="solid")
    cell_font = Font(name="微软雅黑", size=9)
    border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    task_headers = [
        "序号", "任务编号", "项目名称", "检查标准", "检查日期", "任务状态",
        "项目总分", "已评分模块数", "检查项数",
    ]
    module_headers = [
        "任务编号", "项目名称", "检查日期", "检查标准", "模块名称", "模块角色",
        "模块权重", "模块满分", "模块得分", "项目贡献", "检查项数", "满分项数", "扣分项数",
    ]
    detail_headers = [
        "任务编号", "项目名称", "检查日期", "检查标准", "模块名称", "检查项编号",
        "检查项名称", "检查标准/要求", "检查方法", "评分规则", "单项满分", "权重",
        "AI得分", "加权得分", "置信度", "评分依据", "改进建议", "是否跳过", "问题描述",
    ] + [f"问题照片{i}" for i in range(1, max_photo_count + 1)]

    for ws, headers in ((ws_tasks, task_headers), (ws_modules, module_headers), (ws_details, detail_headers)):
        for column, header in enumerate(headers, 1):
            _apply_header(ws.cell(row=1, column=column, value=header), header_font, header_fill, border)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}1"
        ws.row_dimensions[1].height = 28

    template_cache = {}
    task_row = 2
    module_row = 2
    detail_row = 2

    for task_index, task in enumerate(tasks, 1):
        task_results = results_by_task.get(task.task_id, [])
        if not task_results:
            continue
        standard_type = task.standard_type or "diecheng"
        is_lizhi = standard_type == "lizhi"
        module_results = defaultdict(list)
        aggregate_input = {}
        for result in task_results:
            module_results[result.module_name].append(result)
        for module_name, rows in module_results.items():
            aggregate_input[module_name] = {
                "raw_score_sum": sum(float(row.weighted_score or 0) for row in rows),
                "weight_sum": sum(float(row.weight or 0) for row in rows),
            }
        aggregation = aggregate(standard_type, aggregate_input)

        project_name = task.project.name if task.project else ""
        _write_row(
            ws_tasks,
            task_row,
            [
                task_index, task.task_id, project_name, STANDARD_LABELS.get(standard_type, standard_type),
                task.check_date, STATUS_LABELS.get(task.status, task.status),
                float(task.total_score if task.total_score is not None else aggregation["total"]),
                len(module_results), len(task_results),
            ],
            cell_font,
            border,
            center_columns=(1, 4, 5, 6, 7, 8, 9),
        )
        task_row += 1

        configured_modules = standards.get_modules(standard_type)
        ordered_modules = configured_modules + [name for name in module_results if name not in configured_modules]
        for module_name in ordered_modules:
            rows = module_results.get(module_name, [])
            if not rows:
                continue
            cfg = standards.get_module_cfg(standard_type, module_name)
            module_score = float(aggregation["module_pct"].get(module_name, 0))
            contribution = module_score if is_lizhi else module_score * float(cfg.get("weight", 0) or 0)
            role = cfg.get("role", "score")
            full_count = 0
            deduct_count = 0
            for row in rows:
                if row.is_skipped:
                    continue
                maximum = _item_max_score(standard_type, module_name, row.item_id, template_cache)
                is_full = float(row.score or 0) >= (0 if role == "deduction" else maximum)
                full_count += int(is_full)
                deduct_count += int(not is_full)
            _write_row(
                ws_modules,
                module_row,
                [
                    task.task_id, project_name, task.check_date, STANDARD_LABELS.get(standard_type, standard_type),
                    module_name, "扣分模块" if role == "deduction" else "得分模块",
                    float(cfg.get("weight", 0) or 0), cfg.get("max_score"), round(module_score, 2),
                    round(contribution, 2), len(rows), full_count, deduct_count,
                ],
                cell_font,
                border,
                center_columns=(3, 4, 6, 7, 8, 9, 10, 11, 12, 13),
            )
            module_row += 1

            for result in rows:
                key = (result.record_id, result.module_name, result.item_id)
                item_issues = issues_by_item.get(key, [])
                issue_texts = []
                for issue in item_issues:
                    text = issue.description or ""
                    if issue.severity:
                        text += f" [{issue.severity}]"
                    if issue.location:
                        text += f" @ {issue.location}"
                    issue_texts.append(text)

                values = [
                    task.task_id, project_name, task.check_date, STANDARD_LABELS.get(standard_type, standard_type),
                    module_name, result.item_id, result.item_name or "", result.check_standard or "",
                    result.check_method or "", result.scoring_rule or "",
                    _item_max_score(standard_type, module_name, result.item_id, template_cache),
                    float(result.weight or 0), float(result.score or 0), float(result.weighted_score or 0),
                    float(result.confidence_score) if result.confidence_score is not None else None,
                    result.scoring_basis or "", result.improvement_suggestion or "",
                    "是" if result.is_skipped else "否", "\n".join(issue_texts),
                ] + [None] * max_photo_count
                _write_row(
                    ws_details,
                    detail_row,
                    values,
                    cell_font,
                    border,
                    center_columns=(3, 4, 11, 12, 13, 14, 15, 18),
                )

                row_has_image = False
                for photo_index, photo in enumerate(photos_by_item.get(key, []), 1):
                    column = len(detail_headers) - max_photo_count + photo_index
                    cell = ws_details.cell(row=detail_row, column=column)
                    path = resolve_photo_path(photo.file_path)
                    if not path:
                        cell.value = f"图片缺失：{photo.file_name or photo.photo_id}"
                        continue
                    try:
                        # Embed a real thumbnail rather than the original file.
                        # Merely changing XlImage.width/height keeps full-resolution
                        # bytes in the workbook and makes all-project exports huge.
                        with PilImage.open(path) as source:
                            source = ImageOps.exif_transpose(source)
                            source.thumbnail((120, 90))
                            if source.mode not in ("RGB", "L"):
                                rgba = source.convert("RGBA")
                                background = PilImage.new("RGB", rgba.size, "white")
                                background.paste(rgba, mask=rgba.getchannel("A"))
                                source = background
                            elif source.mode == "L":
                                source = source.convert("RGB")
                            stream = io.BytesIO()
                            source.save(stream, format="JPEG", quality=78, optimize=True)
                            stream.seek(0)
                            width, height = source.size
                        image_streams.append(stream)
                        image = XlImage(stream)
                        image.width = width
                        image.height = height
                        ws_details.add_image(image, cell.coordinate)
                        row_has_image = True
                    except Exception:
                        cell.value = f"图片读取失败：{photo.file_name or photo.photo_id}"
                if row_has_image:
                    ws_details.row_dimensions[detail_row].height = 72
                detail_row += 1

    for column, width in enumerate([8, 24, 18, 12, 12, 12, 12, 14, 12], 1):
        ws_tasks.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    for column, width in enumerate([24, 18, 12, 12, 18, 12, 12, 12, 12, 12, 12, 12, 12], 1):
        ws_modules.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    detail_widths = [24, 18, 12, 12, 18, 16, 22, 32, 24, 28, 12, 10, 10, 12, 10, 32, 28, 10, 36]
    for column, width in enumerate(detail_widths, 1):
        ws_details.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    for column in range(len(detail_widths) + 1, len(detail_headers) + 1):
        ws_details.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 20

    temp_file = tempfile.NamedTemporaryFile(prefix="ai_scoring_batch_", suffix=".xlsx", delete=False)
    temp_path = temp_file.name
    temp_file.close()
    try:
        wb.save(temp_path)
    except Exception:
        if os.path.exists(temp_path):
            os.remove(temp_path)
        raise
    finally:
        wb.close()
        for stream in image_streams:
            stream.close()
    return temp_path
