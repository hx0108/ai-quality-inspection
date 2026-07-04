"""
整改复查 API
项目人员提交整改 → AI核查 → 管理员审核
"""
import os
import json
import uuid as _uuid
import threading
import asyncio
from datetime import datetime
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from pydantic import BaseModel
from sqlalchemy.orm import Session

from database import get_db
from models.models import (
    User, Issue, Photo, InspectionRecord, InspectionTask,
    Rectification, Project
)
from api.deps import get_current_user, check_role, get_user_project_filter
from config import settings
from core.logger import get_logger

logger = get_logger("rectification")

router = APIRouter()


def _gen_id(prefix: str) -> str:
    """生成唯一ID"""
    today = datetime.now().strftime("%Y%m%d")
    short = _uuid.uuid4().hex[:8]
    return f"{prefix}-{today}-{short}"


# ==================== 请求模型 ====================
class SubmitRequest(BaseModel):
    rectification_note: str = ""
    rectification_time: str = ""  # YYYY-MM-DD HH:mm


class ReviewRequest(BaseModel):
    approved: bool
    review_note: str = ""


class AppealRequest(BaseModel):
    appeal_note: str


# ==================== API 端点 ====================
@router.get("/pending", summary="获取待整改列表（项目人员）")
async def get_pending_rectifications(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """获取当前项目人员所有整改记录（含 AI 审核结果，按项目隔离）"""
    project_filter = get_user_project_filter(current_user, db)
    if project_filter is not None and len(project_filter) == 0:
        return {"total": 0, "items": [], "projects": [], "modules": []}

    # 关联 InspectionTask 和 Project 获取 project_id 过滤
    query = db.query(Rectification).join(InspectionTask)

    if project_filter is not None:
        query = query.filter(InspectionTask.project_id.in_(project_filter))

    records = query.order_by(Rectification.created_at.desc()).all()

    results = []
    for r in records:
        issue = r.issue
        if not issue:
            continue

        # 获取原问题照片
        issue_photos = db.query(Photo).filter(
            Photo.issue_id == issue.issue_id,
            Photo.photo_type == "问题照片"
        ).all()

        # 获取整改照片（如果有）
        rect_photos = db.query(Photo).filter(
            Photo.issue_id == issue.issue_id,
            Photo.photo_type == "整改照片"
        ).all()

        # 获取项目名和任务信息
        record = r.record
        task = r.task
        project_name = task.project.name if task and task.project else ""
        project_address = task.project.address if task and task.project else ""

        results.append({
            "rectification_id": r.rectification_id,
            "issue_id": issue.issue_id,
            "item_id": issue.item_id,
            "item_name": issue.item_name,
            "description": issue.description,
            "severity": issue.severity,
            "location": issue.location,
            "status": r.status,
            "module_name": issue.module_name,
            "project_name": project_name,
            "project_address": project_address,
            "issue_photos": [
                {"photo_id": p.photo_id, "file_name": p.file_name}
                for p in issue_photos
            ],
            "rectification_photos": [
                {"photo_id": p.photo_id, "file_name": p.file_name}
                for p in rect_photos
            ],
            "rectification_note": r.rectification_note or "",
            "rectification_time": r.rectification_time or "",
            "review_note": r.review_note or "",
            "ai_result": json.loads(r.ai_result) if r.ai_result else None,
            "ai_checked_at": r.ai_checked_at.isoformat() if r.ai_checked_at else None,
            "created_at": r.created_at.isoformat(),
            "deadline": r.deadline or ""
        })

    # 收集所有项目名和模块名（用于前端筛选下拉框）
    projects = sorted(set(r.get("project_name", "") for r in results if r.get("project_name")))
    modules = sorted(set(r.get("module_name", "") for r in results if r.get("module_name")))

    return {"total": len(results), "items": results, "projects": projects, "modules": modules}


@router.get("/task/{task_id}", summary="获取任务下所有整改记录")
async def get_task_rectifications(
    task_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """获取某任务下所有整改记录"""
    records = db.query(Rectification).filter(
        Rectification.task_id == task_id
    ).all()

    results = []
    for r in records:
        issue = r.issue
        ai_result = json.loads(r.ai_result) if r.ai_result else None

        results.append({
            "rectification_id": r.rectification_id,
            "issue_id": r.issue_id,
            "item_id": issue.item_id if issue else "",
            "item_name": issue.item_name if issue else "",
            "description": issue.description if issue else "",
            "severity": issue.severity if issue else "",
            "location": issue.location if issue else "",
            "status": r.status,
            "module_name": issue.module_name if issue else "",
            "rectification_note": r.rectification_note or "",
            "rectification_time": r.rectification_time or "",
            "ai_result": ai_result,
            "ai_checked_at": r.ai_checked_at.isoformat() if r.ai_checked_at else None,
            "review_note": r.review_note or "",
            "reviewer_id": r.reviewer_id,
            "reviewed_at": r.reviewed_at.isoformat() if r.reviewed_at else None,
            "created_at": r.created_at.isoformat(),
            "updated_at": r.updated_at.isoformat()
        })

    return {"total": len(results), "items": results}


@router.get("/all", summary="获取所有整改记录")
async def get_all_rectifications(
    status_filter: Optional[str] = None,
    project_name: Optional[str] = None,
    module_name: Optional[str] = None,
    keyword: Optional[str] = None,
    page: int = 1,
    page_size: int = 0,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """获取所有整改记录，支持按状态/项目/模块/关键词筛选（按项目隔离）"""
    project_filter = get_user_project_filter(current_user, db)
    if project_filter is not None and len(project_filter) == 0:
        return {"total": 0, "items": [], "projects": [], "modules": []}

    query = db.query(Rectification).join(InspectionTask)

    if status_filter:
        statuses = status_filter.split(",")
        query = query.filter(Rectification.status.in_(statuses))

    # 项目级隔离：site_supervisor/field_supervisor/project_staff 只能看关联项目
    if project_filter is not None:
        query = query.filter(InspectionTask.project_id.in_(project_filter))

    # 按项目名筛选（仅管理员/检查员可用，field角色无权指定其他项目）
    if project_name:
        if current_user.role in ("admin", "inspector"):
            proj_ids = [p.id for p in db.query(Project).filter(Project.name == project_name).all()]
            if proj_ids:
                query = query.filter(InspectionTask.project_id.in_(proj_ids))
            else:
                return {"total": 0, "items": [], "projects": [], "modules": []}
        # site_supervisor/field_supervisor/project_staff 无权按其他项目名筛选（已被 project_filter 限制）

    # 按模块筛选
    if module_name:
        issue_ids = [i.issue_id for i in db.query(Issue).filter(Issue.module_name == module_name).all()]
        if issue_ids:
            query = query.filter(Rectification.issue_id.in_(issue_ids))
        else:
            return {"total": 0, "items": [], "projects": [], "modules": []}

    records = query.order_by(Rectification.created_at.desc()).all()

    results = []
    all_projects = set()
    all_modules = set()

    for r in records:
        issue = r.issue
        record = r.record
        task = r.task
        proj_name = task.project.name if task and task.project else ""
        ai_result = json.loads(r.ai_result) if r.ai_result else None

        # 整改照片
        rect_photos = db.query(Photo).filter(
            Photo.issue_id == r.issue_id,
            Photo.photo_type == "整改照片"
        ).all()

        # 原问题照片
        issue_photos = db.query(Photo).filter(
            Photo.issue_id == r.issue_id,
            Photo.photo_type == "问题照片"
        ).all()

        # 审核人姓名
        reviewer_name = ""
        if r.reviewer_id:
            reviewer = db.query(User).filter(User.id == r.reviewer_id).first()
            reviewer_name = reviewer.real_name if reviewer else ""

        mod_name = issue.module_name if issue else ""
        desc = issue.description if issue else ""

        # 关键词过滤
        if keyword:
            kw = keyword.lower()
            if kw not in proj_name.lower() and kw not in mod_name.lower() and kw not in (desc or "").lower() and kw not in (issue.item_name or "").lower() and kw not in r.task_id.lower():
                continue

        all_projects.add(proj_name)
        if mod_name:
            all_modules.add(mod_name)

        results.append({
            "rectification_id": r.rectification_id,
            "issue_id": r.issue_id,
            "item_id": issue.item_id if issue else "",
            "item_name": issue.item_name if issue else "",
            "description": desc,
            "severity": issue.severity if issue else "",
            "location": issue.location if issue else "",
            "status": r.status,
            "module_name": mod_name,
            "project_name": proj_name,
            "task_id": r.task_id,
            "check_date": task.check_date if task else "",
            "rectification_note": r.rectification_note or "",
            "rectification_time": r.rectification_time or "",
            "issue_photos": [
                {"photo_id": p.photo_id, "file_name": p.file_name}
                for p in issue_photos
            ],
            "rectification_photos": [
                {"photo_id": p.photo_id, "file_name": p.file_name}
                for p in rect_photos
            ],
            "ai_result": ai_result,
            "ai_checked_at": r.ai_checked_at.isoformat() if r.ai_checked_at else None,
            "review_note": r.review_note or "",
            "reviewer_name": reviewer_name,
            "reviewed_at": r.reviewed_at.isoformat() if r.reviewed_at else None,
            "submitted_at": r.submitted_at.isoformat() if r.submitted_at else None,
            "created_at": r.created_at.isoformat(),
            "deadline": r.deadline or ""
        })

    total = len(results)
    # 分页：page_size>0 时返回对应页（移动端 van-list 用），默认0返回全部（PC 行为不变）
    if page_size and page_size > 0:
        start = (max(page, 1) - 1) * page_size
        paged = results[start:start + page_size]
    else:
        paged = results

    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": paged,
    }


@router.get("/export/excel", summary="导出整改记录Excel")
async def export_rectifications_excel(
    status_filter: Optional[str] = None,
    project_name: Optional[str] = None,
    module_name: Optional[str] = None,
    keyword: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """导出整改记录为Excel文件（含高清照片）"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse
    from PIL import Image as PilImage
    from collections import defaultdict
    from concurrent.futures import ThreadPoolExecutor
    import io

    MAX_PHOTOS = 5
    MAX_DISPLAY_W = 300
    PHOTO_MAX_W = 600

    def compress_photo(photo):
        """加载照片并压缩为 JPEG，返回可安全复用的图片字节及尺寸。"""
        if not os.path.exists(photo.file_path):
            return None
        try:
            img = PilImage.open(photo.file_path)
            if img.mode in ('RGBA', 'P'):
                img = img.convert('RGB')
            w, h = img.size
            if w > PHOTO_MAX_W:
                ratio = PHOTO_MAX_W / w
                w, h = PHOTO_MAX_W, int(h * ratio)
                img = img.resize((w, h), PilImage.LANCZOS)
            buf = io.BytesIO()
            img.save(buf, format='JPEG', quality=75)
            return (photo.photo_id, buf.getvalue(), w, h)
        except Exception:
            return None

    # 项目级隔离
    project_filter = get_user_project_filter(current_user, db)
    if project_filter is not None and len(project_filter) == 0:
        return StreamingResponse(
            io.BytesIO(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=empty.xlsx"}
        )

    # 查询记录
    query = db.query(Rectification).join(InspectionTask)

    if project_filter is not None:
        query = query.filter(InspectionTask.project_id.in_(project_filter))
    if status_filter:
        statuses = status_filter.split(",")
        query = query.filter(Rectification.status.in_(statuses))
    if project_name:
        if current_user.role in ("admin", "inspector"):
            proj_ids = [p.id for p in db.query(Project).filter(Project.name == project_name).all()]
            if proj_ids:
                query = query.filter(InspectionTask.project_id.in_(proj_ids))
            else:
                query = query.filter(Rectification.rectification_id == "__none__")
    if module_name:
        issue_ids = [i.issue_id for i in db.query(Issue).filter(Issue.module_name == module_name).all()]
        if issue_ids:
            query = query.filter(Rectification.issue_id.in_(issue_ids))
        else:
            query = query.filter(Rectification.rectification_id == "__none__")

    records = query.order_by(Rectification.created_at.desc()).all()

    # 预过滤（关键词）并收集 issue_id
    filtered = []
    for r in records:
        issue = r.issue
        task = r.task
        proj_name = task.project.name if task and task.project else ""
        mod_name = issue.module_name if issue else ""
        desc = issue.description if issue else ""
        item_name = issue.item_name if issue else ""
        if keyword:
            kw = keyword.lower()
            if kw not in proj_name.lower() and kw not in mod_name.lower() and kw not in desc.lower() and kw not in item_name.lower() and kw not in r.task_id.lower():
                continue
        filtered.append(r)

    # 批量查询所有照片（一次查询代替 N×2 次）
    all_issue_ids = [r.issue.issue_id for r in filtered if r.issue]
    photo_lookup = defaultdict(list)  # (issue_id, photo_type) -> [Photo]
    if all_issue_ids:
        all_photos = db.query(Photo).filter(
            Photo.issue_id.in_(all_issue_ids),
            Photo.photo_type.in_(["问题照片", "整改照片"])
        ).all()
        for p in all_photos:
            photo_lookup[(p.issue_id, p.photo_type)].append(p)
        # 每类只取前 MAX_PHOTOS 张
        for key in photo_lookup:
            photo_lookup[key] = photo_lookup[key][:MAX_PHOTOS]

    # 并行压缩所有照片
    flat_photos = [p for plist in photo_lookup.values() for p in plist]
    photo_buffers = {}  # photo_id -> (bytes, w, h)
    if flat_photos:
        with ThreadPoolExecutor(max_workers=4) as pool:
            for result in pool.map(compress_photo, flat_photos):
                if result:
                    photo_buffers[result[0]] = (result[1], result[2], result[3])

    # 构建Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "整改记录"

    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    headers = [
        "序号", "项目", "任务ID", "模块", "检查项", "问题描述", "严重程度",
        "位置", "整改说明", "整改时间", "整改状态",
        "AI评分", "AI建议", "AI分析", "审核意见", "审核人", "创建时间",
        "检查照片1", "检查照片2", "检查照片3", "检查照片4", "检查照片5",
        "整改照片1", "整改照片2", "整改照片3", "整改照片4", "整改照片5"
    ]
    ws.append(headers)

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    status_map = {
        "pending": "待整改", "submitted": "AI核查中",
        "ai_approved": "AI通过", "ai_rejected": "AI驳回",
        "approved": "已通过"
    }
    ws.row_dimensions[1].height = 30

    for idx, r in enumerate(filtered, 1):
        issue = r.issue
        task = r.task
        proj_name = task.project.name if task and task.project else ""
        ai_result = json.loads(r.ai_result) if r.ai_result else None

        reviewer_name = ""
        if r.reviewer_id:
            reviewer = db.query(User).filter(User.id == r.reviewer_id).first()
            reviewer_name = reviewer.real_name if reviewer else ""

        row_data = [
            idx,
            proj_name,
            r.task_id,
            issue.module_name if issue else "",
            issue.item_name if issue else "",
            issue.description if issue else "",
            issue.severity if issue else "",
            issue.location if issue else "",
            r.rectification_note or "",
            r.rectification_time or "",
            status_map.get(r.status, r.status),
            ai_result.get("confidence_score", "") if ai_result else "",
            ai_result.get("suggestion", "") if ai_result else "",
            ai_result.get("analysis", "") if ai_result else "",
            r.review_note or "",
            reviewer_name,
            r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else ""
        ] + [""] * (MAX_PHOTOS * 2)
        ws.append(row_data)

        row_num = idx + 1
        max_img_h = 0

        if issue:
            photo_col_starts = {"问题照片": 18, "整改照片": 18 + MAX_PHOTOS}
            for ptype, start_col in photo_col_starts.items():
                photos = photo_lookup.get((issue.issue_id, ptype), [])
                for pi, photo in enumerate(photos):
                    cached = photo_buffers.get(photo.photo_id)
                    if not cached:
                        continue
                    image_data, orig_w, orig_h = cached
                    try:
                        # A workbook can contain the same photo in multiple
                        # rectification rows; openpyxl closes each image stream
                        # while saving, so every placement needs its own stream.
                        xl_img = XlImage(io.BytesIO(image_data))
                        if orig_w > MAX_DISPLAY_W:
                            scale = MAX_DISPLAY_W / orig_w
                            xl_img.width = int(orig_w * scale)
                            xl_img.height = int(orig_h * scale)
                        else:
                            xl_img.width = orig_w
                            xl_img.height = orig_h
                        max_img_h = max(max_img_h, xl_img.height)
                        col_letter = get_column_letter(start_col + pi)
                        ws.add_image(xl_img, f'{col_letter}{row_num}')
                    except Exception:
                        pass

        if max_img_h > 0:
            ws.row_dimensions[row_num].height = max_img_h * 0.75 + 6

    # 设置列宽
    photo_col_w = 28
    col_widths = [6, 15, 20, 12, 16, 30, 10, 12, 30, 18, 10, 8, 8, 30, 20, 10, 18] + [photo_col_w] * (MAX_PHOTOS * 2)
    for i, w in enumerate(col_widths, 1):
        col_letter = chr(64 + i) if i <= 26 else chr(64 + i // 26) + chr(65 + i % 26)
        ws.column_dimensions[col_letter].width = w

    # 数据行样式
    data_align = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.alignment = data_align
            cell.border = thin_border

    # 输出到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    from urllib.parse import quote
    filename = f"rectifications_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    cn_filename = quote(f"整改记录_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{filename}\"; filename*=UTF-8''{cn_filename}"}
    )


@router.get("/{rectification_id}", summary="获取整改详情")
async def get_rectification_detail(
    rectification_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """获取整改详情（含AI结果和照片，按项目隔离）"""
    r = db.query(Rectification).filter(
        Rectification.rectification_id == rectification_id
    ).first()
    if not r:
        raise HTTPException(status_code=404, detail="整改记录不存在")

    # 项目级隔离检查
    project_filter = get_user_project_filter(current_user, db)
    if project_filter is not None and len(project_filter) > 0:
        task = db.query(InspectionTask).filter(InspectionTask.task_id == r.task_id).first()
        if not task or task.project_id not in project_filter:
            raise HTTPException(status_code=403, detail="无权访问此整改记录")

    issue = r.issue
    task = r.task
    project_name = task.project.name if task and task.project else ""
    project_address = task.project.address if task and task.project else ""
    ai_result = json.loads(r.ai_result) if r.ai_result else None

    # 原问题照片
    issue_photos = db.query(Photo).filter(
        Photo.issue_id == r.issue_id,
        Photo.photo_type == "问题照片"
    ).all()

    # 整改照片
    rect_photos = db.query(Photo).filter(
        Photo.issue_id == r.issue_id,
        Photo.photo_type == "整改照片"
    ).all()

    return {
        "rectification_id": r.rectification_id,
        "issue_id": r.issue_id,
        "record_id": r.record_id,
        "task_id": r.task_id,
        "status": r.status,
        "item_id": issue.item_id if issue else "",
        "item_name": issue.item_name if issue else "",
        "description": issue.description if issue else "",
        "severity": issue.severity if issue else "",
        "location": issue.location if issue else "",
        "module_name": issue.module_name if issue else "",
        "project_name": project_name,
        "project_address": project_address,
        "check_date": task.check_date if task else "",
        "rectification_note": r.rectification_note or "",
        "rectification_time": r.rectification_time or "",
        "issue_photos": [
            {"photo_id": p.photo_id, "file_name": p.file_name}
            for p in issue_photos
        ],
        "rectification_photos": [
            {"photo_id": p.photo_id, "file_name": p.file_name}
            for p in rect_photos
        ],
        "ai_result": ai_result,
        "ai_checked_at": r.ai_checked_at.isoformat() if r.ai_checked_at else None,
        "review_note": r.review_note or "",
        "reviewed_at": r.reviewed_at.isoformat() if r.reviewed_at else None,
        "submitted_at": r.submitted_at.isoformat() if r.submitted_at else None,
        "created_at": r.created_at.isoformat(),
        "updated_at": r.updated_at.isoformat()
    }


@router.post("/{rectification_id}/submit", summary="提交整改")
async def submit_rectification(
    rectification_id: str,
    request: SubmitRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(check_role(["project_staff", "field_supervisor"]))
):
    """
    项目人员提交整改材料 → 触发AI核查
    - AI通过(≥85分) → 自动标记 ai_approved
    - AI驳回(<85分) → 标记 ai_rejected，等待管理员审核
    """
    logger.info(f"提交整改请求: rectification_id={rectification_id}, user={current_user.username}, data={request}")

    try:
        r = db.query(Rectification).filter(
            Rectification.rectification_id == rectification_id
        ).first()
        if not r:
            logger.warning(f"整改记录不存在: {rectification_id}")
            raise HTTPException(status_code=404, detail="整改记录不存在")

        # 项目级隔离检查
        project_filter = get_user_project_filter(current_user, db)
        if project_filter is not None and len(project_filter) > 0:
            task = db.query(InspectionTask).filter(InspectionTask.task_id == r.task_id).first()
            if not task or task.project_id not in project_filter:
                logger.warning(f"项目权限检查失败: user={current_user.username}, project_filter={project_filter}, task.project_id={task.project_id if task else 'N/A'}")
                raise HTTPException(status_code=403, detail="无权访问此整改记录")

        logger.info(f"当前状态: {r.status}")
        if r.status not in ["pending"]:
            logger.warning(f"状态不允许提交: {r.status}")
            raise HTTPException(status_code=400, detail="当前状态不允许提交")

        # 检查是否有整改照片
        rect_photos = db.query(Photo).filter(
            Photo.issue_id == r.issue_id,
            Photo.photo_type == "整改照片"
        ).all()
        logger.info(f"整改照片数量: {len(rect_photos)}, issue_id={r.issue_id}")

        if not rect_photos:
            raise HTTPException(status_code=400, detail="请先上传至少一张整改照片")

        # 更新整改信息
        r.rectification_note = request.rectification_note
        r.rectification_time = request.rectification_time
        r.submitted_at = datetime.utcnow()
        r.status = "submitted"
        db.commit()
        logger.info(f"整改提交成功: {rectification_id}")

        # 通知管理员有新的整改提交
        try:
            from api.notification import send_notification
            task = db.query(InspectionTask).filter(InspectionTask.task_id == r.task_id).first()
            project_name = ""
            if task:
                from models.models import Project
                proj = db.query(Project).filter(Project.id == task.project_id).first()
                project_name = proj.name if proj else ""
            admins = db.query(User).filter(User.role == 'admin', User.is_active == True).all()
            for admin in admins:
                send_notification(
                    user_id=admin.id, username=admin.username,
                    title=f"整改已提交 - {project_name}",
                    content=f"项目「{project_name}」有新的整改已提交，请及时审核。",
                    notify_type="rectification_reminder", ref_type="rectification",
                    ref_id=rectification_id, async_send=True
                )
        except Exception as ne:
            logger.warning(f"整改提交通知发送失败: {ne}")
        _trigger_ai_check(r.rectification_id)

        return {
            "message": "整改已提交，AI正在核查中",
            "rectification_id": rectification_id,
            "status": "submitted"
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"提交整改异常: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"提交失败: {str(e)}")


@router.post("/{rectification_id}/photos", summary="上传整改照片")
async def upload_rectification_photo(
    rectification_id: str,
    file: UploadFile = File(...),
    metadata: str = Form(""),
    db: Session = Depends(get_db),
    current_user: User = Depends(check_role(["project_staff", "field_supervisor"]))
):
    """上传整改照片（支持水印元数据）"""
    r = db.query(Rectification).filter(
        Rectification.rectification_id == rectification_id
    ).first()
    if not r:
        raise HTTPException(status_code=404, detail="整改记录不存在")

    # 项目级隔离检查
    project_filter = get_user_project_filter(current_user, db)
    if project_filter is not None and len(project_filter) > 0:
        task = db.query(InspectionTask).filter(InspectionTask.task_id == r.task_id).first()
        if not task or task.project_id not in project_filter:
            raise HTTPException(status_code=403, detail="无权访问此整改记录")

    # 验证文件类型
    if not file.filename.lower().endswith(('.jpg', '.jpeg', '.png', '.gif')):
        raise HTTPException(status_code=400, detail="只支持 jpg/png/gif 格式")

    # 解析水印元数据
    photo_meta = {}
    if metadata:
        try:
            photo_meta = json.loads(metadata)
        except (json.JSONDecodeError, TypeError):
            photo_meta = {}

    # 生成文件名
    file_ext = os.path.splitext(file.filename)[1]
    photo_id = _gen_id("PHO")
    filename = f"{photo_id}{file_ext}"

    # 创建保存目录（与检查记录同一目录）
    save_dir = settings.PHOTOS_DIR / r.record_id
    save_dir.mkdir(parents=True, exist_ok=True)

    # 保存文件
    file_path = save_dir / filename
    content = await file.read()
    with open(file_path, "wb") as f:
        f.write(content)

    # 解析拍摄时间
    capture_time = None
    if photo_meta.get("capture_time"):
        try:
            capture_time = datetime.fromisoformat(photo_meta["capture_time"])
        except (ValueError, TypeError):
            capture_time = None

    # 创建照片记录（含水印元数据）
    photo = Photo(
        photo_id=photo_id,
        issue_id=r.issue_id,
        file_path=str(file_path),
        file_name=filename,
        photo_type="整改照片",
        latitude=photo_meta.get("latitude"),
        longitude=photo_meta.get("longitude"),
        address=photo_meta.get("address"),
        capture_time=capture_time,
        inspector_name=photo_meta.get("inspector_name"),
        watermark_hash=photo_meta.get("watermark_hash"),
    )
    db.add(photo)
    db.commit()

    return {
        "message": "照片上传成功",
        "photo_id": photo.photo_id,
        "file_name": filename
    }


@router.delete("/{rectification_id}/photos/{photo_id}", summary="删除整改照片")
async def delete_rectification_photo(
    rectification_id: str,
    photo_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """删除整改照片"""
    photo = db.query(Photo).filter(
        Photo.photo_id == photo_id,
        Photo.photo_type == "整改照片"
    ).first()
    if not photo:
        raise HTTPException(status_code=404, detail="照片不存在")

    r = db.query(Rectification).filter(
        Rectification.rectification_id == rectification_id
    ).first()
    if not r or photo.issue_id != r.issue_id:
        raise HTTPException(status_code=400, detail="照片不属于该整改记录")

    # 项目级隔离检查
    project_filter = get_user_project_filter(current_user, db)
    if project_filter is not None and len(project_filter) > 0:
        task = db.query(InspectionTask).filter(InspectionTask.task_id == r.task_id).first()
        if not task or task.project_id not in project_filter:
            raise HTTPException(status_code=403, detail="无权访问此整改记录")

    # 删除文件
    if os.path.exists(photo.file_path):
        os.remove(photo.file_path)

    db.delete(photo)
    db.commit()

    return {"message": "照片已删除"}


@router.put("/{rectification_id}/review", summary="管理员审核整改")
async def review_rectification(
    rectification_id: str,
    request: ReviewRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(check_role(["admin", "inspector", "site_supervisor"]))
):
    """管理员/检查人员/阵地督导审核整改结果"""
    r = db.query(Rectification).filter(
        Rectification.rectification_id == rectification_id
    ).first()
    if not r:
        raise HTTPException(status_code=404, detail="整改记录不存在")

    if r.status not in ["ai_rejected", "submitted", "ai_approved", "pending_review"]:
        raise HTTPException(status_code=400, detail="当前状态不允许审核")

    if request.approved:
        r.status = "approved"
    else:
        # 驳回后重置为 pending，允许项目人员重新上传整改
        r.status = "pending"
        # 清除旧的整改照片，确保下次AI核查只比对新上传的照片
        old_photos = db.query(Photo).filter(
            Photo.issue_id == r.issue_id,
            Photo.photo_type == "整改照片"
        ).all()
        for photo in old_photos:
            if os.path.exists(photo.file_path):
                os.remove(photo.file_path)
            db.delete(photo)
        # 清除上一次AI核查结果
        r.ai_result = None
        r.ai_checked_at = None
        r.rectification_note = None
        r.rectification_time = None
        # 递增整改轮次
        r.round_number = getattr(r, 'round_number', 1) or 1
        r.round_number += 1

    r.review_note = request.review_note
    r.reviewer_id = current_user.id
    r.reviewed_at = datetime.utcnow()
    r.updated_at = datetime.utcnow()
    db.commit()

    # 通知相关人员：审核结果
    try:
        from api.notification import send_notification
        task = db.query(InspectionTask).filter(InspectionTask.task_id == r.task_id).first()
        project_name = ""
        project_id = None
        if task:
            from models.models import Project
            proj = db.query(Project).filter(Project.id == task.project_id).first()
            project_name = proj.name if proj else ""
            project_id = task.project_id

        # 获取项目关联的驻场经理
        recipients = set()
        if project_id:
            from models.models import UserProject
            ups = db.query(UserProject).filter(UserProject.project_id == project_id).all()
            for up in ups:
                u = db.query(User).filter(User.id == up.user_id, User.role == 'field_supervisor', User.is_active == True).first()
                if u:
                    recipients.add((u.id, u.username))
        # 也通知所有管理员
        for admin in db.query(User).filter(User.role == 'admin', User.is_active == True).all():
            recipients.add((admin.id, admin.username))

        notify_type = "rectification_approved" if request.approved else "rectification_rejected"
        action_text = "已通过" if request.approved else "已驳回"
        for uid, uname in recipients:
            send_notification(
                user_id=uid, username=uname,
                title=f"整改{action_text} - {project_name}",
                content=f"项目「{project_name}」的整改记录已由{current_user.real_name or current_user.username}{action_text}。",
                notify_type=notify_type, ref_type="rectification",
                ref_id=rectification_id, async_send=True
            )
    except Exception as ne:
        logger.warning(f"整改审核通知发送失败: {ne}")

    return {
        "message": "审核完成",
        "rectification_id": rectification_id,
        "status": r.status
    }


@router.post("/{rectification_id}/recheck", summary="重新AI核查")
async def recheck_rectification(
    rectification_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(check_role(["admin", "inspector", "site_supervisor"]))
):
    """管理员/检查人员/阵地督导手动触发重新AI核查"""
    r = db.query(Rectification).filter(
        Rectification.rectification_id == rectification_id
    ).first()
    if not r:
        raise HTTPException(status_code=404, detail="整改记录不存在")

    # 仅重置AI结果，保留现有整改照片供AI重新核查
    r.status = "submitted"
    r.ai_result = None
    r.ai_checked_at = None
    db.commit()

    _trigger_ai_check(r.rectification_id)

    return {"message": "已触发重新AI核查"}


@router.post("/{rectification_id}/appeal", summary="项目人员申诉")
async def appeal_rectification(
    rectification_id: str,
    request: AppealRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(check_role(["project_staff", "field_supervisor"]))
):
    """项目人员/驻场经理对AI驳回结果提交申诉，状态转为待人工审核"""
    r = db.query(Rectification).filter(
        Rectification.rectification_id == rectification_id
    ).first()
    if not r:
        raise HTTPException(status_code=404, detail="整改记录不存在")

    # 项目级隔离检查
    project_filter = get_user_project_filter(current_user, db)
    if project_filter is not None and len(project_filter) > 0:
        task = db.query(InspectionTask).filter(InspectionTask.task_id == r.task_id).first()
        if not task or task.project_id not in project_filter:
            raise HTTPException(status_code=403, detail="无权访问此整改记录")

    if r.status != "ai_rejected":
        raise HTTPException(status_code=400, detail="仅AI驳回状态可申诉")

    r.status = "pending_review"
    r.review_note = f"[申诉] {request.appeal_note}"
    r.updated_at = datetime.utcnow()
    db.commit()

    return {
        "message": "申诉已提交，等待管理员审核",
        "rectification_id": rectification_id,
        "status": "pending_review"
    }


# ==================== AI 核查后台任务（LangGraph） ====================
def _trigger_ai_check(rectification_id: str):
    """后台线程触发 AI 核查 — 使用 LangGraph 工作流"""
    def _run():
        from database import SessionLocal
        db = SessionLocal()
        try:
            r = db.query(Rectification).filter(
                Rectification.rectification_id == rectification_id
            ).first()
            if not r:
                return

            issue = r.issue
            if not issue:
                return

            # 获取原问题照片路径
            issue_photos = db.query(Photo).filter(
                Photo.issue_id == issue.issue_id,
                Photo.photo_type == "问题照片"
            ).all()
            issue_photo_paths = [p.file_path for p in issue_photos if os.path.exists(p.file_path)]

            # 获取整改照片路径
            rect_photos = db.query(Photo).filter(
                Photo.issue_id == issue.issue_id,
                Photo.photo_type == "整改照片"
            ).all()
            rect_photo_paths = [p.file_path for p in rect_photos if os.path.exists(p.file_path)]

            if not rect_photo_paths:
                r.ai_result = json.dumps({
                    "photo_valid": False,
                    "photo_info": "未找到整改照片",
                    "note_qualified": False,
                    "note_info": "",
                    "confidence_score": 0,
                    "rectification_depth": "无法判断",
                    "analysis": "未找到整改照片",
                    "suggestion": "驳回"
                }, ensure_ascii=False)
                r.status = "ai_rejected"
                r.ai_checked_at = datetime.utcnow()
                db.commit()
                return

            # 获取预期地点
            task = r.task
            expected_location = ""
            if task and task.project:
                expected_location = task.project.name

            # 获取当前轮次
            round_number = getattr(r, 'round_number', 1) or 1

            # 构建输入状态
            input_state = {
                "rectification_id": rectification_id,
                "issue_id": r.issue_id,
                "record_id": r.record_id,
                "task_id": r.task_id,
                "issue_description": issue.description or "",
                "issue_severity": issue.severity or "一般",
                "issue_location": issue.location or "",
                "issue_module_name": issue.module_name or "",
                "expected_location": expected_location,
                "issue_photo_paths": issue_photo_paths,
                "rectification_photo_paths": rect_photo_paths,
                "rectification_note": r.rectification_note or "",
                "rectification_time": r.rectification_time or "",
                "round_number": round_number,
                "started_at": datetime.now().isoformat(),
            }

            db.close()
            db = None  # 防止 finally 中再次关闭

            # 运行 LangGraph 工作流
            from agents.rectification.graph import get_rectification_graph
            graph = get_rectification_graph()

            loop = asyncio.new_event_loop()
            try:
                result = loop.run_until_complete(graph.ainvoke(input_state))
                logger.info(f"[LangGraph] 整改验证完成: {rectification_id}, status={result.get('final_status')}")
            finally:
                loop.close()

        except Exception as e:
            logger.error(f"[LangGraph] 整改验证异常: {e}")
            import traceback
            traceback.print_exc()
            # 异常兜底：标记为驳回
            try:
                if db is None:
                    from database import SessionLocal as _sl
                    _db = _sl()
                else:
                    _db = db
                r = _db.query(Rectification).filter(
                    Rectification.rectification_id == rectification_id
                ).first()
                if r:
                    r.status = "ai_rejected"
                    r.ai_result = json.dumps({
                        "photo_valid": False, "photo_info": "",
                        "note_qualified": False, "note_info": "",
                        "confidence_score": 0,
                        "rectification_depth": "无法判断",
                        "analysis": f"AI核查异常: {str(e)}",
                        "suggestion": "驳回"
                    }, ensure_ascii=False)
                    r.ai_checked_at = datetime.utcnow()
                    _db.commit()
                if db is None:
                    _db.close()
            except Exception:
                pass
        finally:
            if db:
                db.close()

    t = threading.Thread(target=_run, daemon=True)
    t.start()
