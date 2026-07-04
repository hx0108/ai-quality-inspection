"""
报告 API
Agent 3: 使用 DeepSeek-V3.2 进行模块分析和报告生成
"""
import json
import os
import time
import uuid as _uuid
from datetime import datetime
from typing import List, Optional
from collections import defaultdict
from fastapi import APIRouter, Depends, HTTPException, status, BackgroundTasks, Request
from fastapi.responses import FileResponse
from jose import jwt
from pydantic import BaseModel
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func

from database import get_db, SessionLocal
from models.models import User, InspectionTask, InspectionRecord, Issue, ScoringResult, Report
from api.deps import get_current_user, check_role, get_user_project_filter
from config import settings
from core.llm_client import DeepSeekClient
from core.logger import get_logger

logger = get_logger("report")

router = APIRouter()

# 报告生成进度缓存（内存+数据库双写）
report_progress = {}  # {task_id: {"status": "generating/completed/failed", "current_step": str, "modules_done": int, "modules_total": int}}

# 报告生成速率限制：防止同一任务频繁触发
_report_trigger_times: dict = {}  # {task_id: last_trigger_timestamp}
_REPORT_RATE_LIMIT = 60  # 60秒内不能重复触发同一任务的报告生成


def _persist_report_progress(task_id: str):
    """将报告进度持久化到数据库"""
    try:
        from core.progress_store import progress_store
        data = report_progress.get(task_id, {})
        progress_store.update(
            "report", task_id,
            status=data.get("status", "generating"),
            detail={
                "current_step": data.get("current_step", ""),
                "modules_done": data.get("modules_done", 0),
                "modules_total": data.get("modules_total", 0),
            },
            error=data.get("error")
        )
    except Exception:
        pass


# ==================== 请求/响应模型 ====================
class ReportGenerate(BaseModel):
    task_id: str
    report_type: str = "full"


# ==================== 报告生成逻辑 ====================
def run_report_generation_sync(task_id: str):
    """同步方式运行报告生成任务（在BackgroundTasks线程中执行）"""
    import asyncio as asyncio_module

    report_progress[task_id] = {
        "status": "generating",
        "current_step": "collecting_data",
        "modules_done": 0,
        "modules_total": 0,
        "error": None
    }
    _persist_report_progress(task_id)

    loop = asyncio_module.new_event_loop()
    asyncio_module.set_event_loop(loop)
    try:
        logger.info(f"开始生成报告: {task_id}")
        loop.run_until_complete(_generate_report_async(task_id))
        report_progress[task_id]["status"] = "completed"
        _persist_report_progress(task_id)
        logger.info(f"报告生成完成: {task_id}")

        # 通知管理员报告已生成
        try:
            from api.notification import send_notification
            task = SessionLocal().query(InspectionTask).filter(InspectionTask.task_id == task_id).first()
            project_name = ""
            if task:
                from models.models import Project
                proj = SessionLocal().query(Project).filter(Project.id == task.project_id).first()
                project_name = proj.name if proj else ""
            db_notif = SessionLocal()
            admins = db_notif.query(User).filter(User.role == 'admin', User.is_active == True).all()
            for admin in admins:
                send_notification(
                    user_id=admin.id,
                    username=admin.username,
                    title=f"报告已生成 - {project_name}",
                    content=f"项目「{project_name}」的检查报告已生成完成，请查阅。",
                    notify_type="report_ready",
                    ref_type="task",
                    ref_id=task_id,
                    async_send=True
                )
            db_notif.close()
        except Exception as ne:
            logger.warning(f"报告完成通知发送失败: {ne}")
    except Exception as e:
        logger.error(f"报告生成失败: {e}")
        import traceback
        traceback.print_exc()
        report_progress[task_id]["status"] = "failed"
        report_progress[task_id]["error"] = str(e)
        _persist_report_progress(task_id)
    finally:
        loop.close()


async def _generate_report_async(task_id: str):
    """异步报告生成核心逻辑 — 使用 LangGraph Agent 3"""
    from agents.report import build_report_graph

    report_progress[task_id]["modules_total"] = 0

    graph = build_report_graph()
    result = await graph.ainvoke({
        "task_id": task_id,
        "module_analyses": [],
        "errors": [],
        "modules_done": 0,
        "modules_total": 0,
    })

    # 更新进度中的模块数
    scored_count = len([m for m in result.get("module_scores", []) if m.get("module_pct_score", 0) > 0])
    report_progress[task_id]["modules_total"] = scored_count
    report_progress[task_id]["modules_done"] = scored_count
    _persist_report_progress(task_id)

    if result.get("errors"):
        raise Exception("; ".join(result["errors"]))


def generate_pdf_report(report_content: dict, output_path: str) -> str:
    """生成 PDF 格式报告（使用 reportlab）"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm, cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    # 注册中文字体（优先外部 TTF，fallback 到 reportlab 内置 CID 字体）
    import platform
    import os
    import glob

    font_name = None
    use_cid = False  # 是否使用内置 CID 字体

    def _try_register_ttf(font_file, name='SimHei'):
        """尝试注册 TTF 字体，成功返回 True"""
        if not os.path.exists(font_file) or os.path.getsize(font_file) == 0:
            return False
        try:
            pdfmetrics.registerFont(TTFont(name, font_file))
            return True
        except Exception:
            return False

    if platform.system() == 'Windows':
        for c in [r'C:\Windows\Fonts\simhei.ttf', r'C:\Windows\Fonts\msyh.ttc']:
            if _try_register_ttf(c):
                pdfmetrics.registerFont(TTFont('SimHei-Bold', c))
                font_name = 'SimHei'
                break
    elif platform.system() == 'Darwin':
        for c in ['/System/Library/Fonts/STHeiti Medium.ttc', '/Library/Fonts/Arial Unicode.ttf']:
            if _try_register_ttf(c):
                pdfmetrics.registerFont(TTFont('SimHei-Bold', c))
                font_name = 'SimHei'
                break
    else:  # Linux / Docker
        # 搜索所有可能的 TTF/TTC 中文字体
        candidates = []
        for pattern in ['/app/fonts/*.ttf', '/app/fonts/*.otf',
                        '/usr/share/fonts/**/*.ttf', '/usr/share/fonts/**/*.ttc']:
            candidates.extend(glob.glob(pattern, recursive=True))
        # 按优先级排序：simhei > wqy > noto > 其他
        def _priority(p):
            p_lower = p.lower()
            for i, kw in enumerate(['simhei', 'wqy', 'noto', 'cjk']):
                if kw in p_lower:
                    return i
            return 99
        candidates.sort(key=_priority)
        for c in candidates:
            if _try_register_ttf(c):
                pdfmetrics.registerFont(TTFont('SimHei-Bold', c))
                font_name = 'SimHei'
                break

    # Fallback: 使用 reportlab 内置 CID 字体（无需外部文件）
    if not font_name:
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        font_name = 'STSong-Light'
        use_cid = True

    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm
    )

    # 样式（使用动态 font_name，兼容外部字体和内置 CID 字体）
    styles = getSampleStyleSheet()
    bold_font = 'SimHei-Bold' if not use_cid else font_name
    style_title = ParagraphStyle('CNTitle', parent=styles['Title'], fontName=font_name, fontSize=20, spaceAfter=20)
    style_h1 = ParagraphStyle('CNH1', parent=styles['Heading1'], fontName=font_name, fontSize=16, spaceAfter=10, spaceBefore=16)
    style_h2 = ParagraphStyle('CNH2', parent=styles['Heading2'], fontName=font_name, fontSize=13, spaceAfter=8, spaceBefore=12)
    style_body = ParagraphStyle('CNBody', parent=styles['Normal'], fontName=font_name, fontSize=10, leading=16, spaceAfter=4)
    style_bullet = ParagraphStyle('CNBullet', parent=style_body, leftIndent=20, bulletIndent=10)
    style_bold = ParagraphStyle('CNBold', parent=style_body, fontName=bold_font)

    elements = []

    # 标题
    elements.append(Paragraph('物业品质检查报告', style_title))
    elements.append(Spacer(1, 10 * mm))

    # 一、检查概况
    elements.append(Paragraph('一、检查概况', style_h1))
    info_data = [
        ['项目名称', report_content.get('project_name', '')],
        ['检查日期', str(report_content.get('check_date', ''))],
        ['项目总分', f"{report_content.get('total_score', 0):.2f} 分"],
        ['问题总数', str(report_content.get('issue_summary', {}).get('total_count', 0))],
        ['评分方法', f"{report_content.get('scoring_method', '')} + {report_content.get('analysis_method', '')}"],
    ]
    info_table = Table(info_data, colWidths=[3.5 * cm, 12 * cm])
    info_table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('BACKGROUND', (0, 0), (0, -1), colors.Color(0.95, 0.95, 0.95)),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 8 * mm))

    # 二、各模块得分
    elements.append(Paragraph('二、各模块检查情况', style_h1))
    modules = report_content.get('modules', [])
    if modules:
        score_data = [['模块名称', '得分', '权重', '贡献分', '问题数']]
        for m in modules:
            if m.get('module_pct_score', 0) > 0:
                score_data.append([
                    m['module_name'],
                    f"{m['module_pct_score']:.2f}",
                    f"{m['weight_ratio'] * 100:.0f}%",
                    f"{m['weighted_contribution']:.2f}",
                    str(m.get('issue_count', 0))
                ])
        score_table = Table(score_data, colWidths=[3.5 * cm, 2.5 * cm, 2 * cm, 2.5 * cm, 2 * cm])
        score_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, -1), font_name),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BACKGROUND', (0, 0), (-1, 0), colors.Color(0.1, 0.4, 0.7)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.Color(0.97, 0.97, 0.97)]),
        ]))
        elements.append(score_table)
    elements.append(Spacer(1, 8 * mm))

    # 三、重点问题清单
    elements.append(Paragraph('三、重点问题清单', style_h1))
    for severity in ['严重', '一般', '轻微']:
        issues = report_content.get('issues', {}).get(severity, [])
        if issues:
            idx = ['严重', '一般', '轻微'].index(severity) + 1
            elements.append(Paragraph(f'3.{idx} {severity}问题（{len(issues)}项）', style_h2))
            for issue in issues[:10]:
                text = f"[{issue.get('module_name', '')}] {issue.get('item_name', '')}"
                if issue.get('description'):
                    text += f" — {issue['description']}"
                if issue.get('location'):
                    text += f" (位置: {issue['location']})"
                elements.append(Paragraph(text, style_bullet))
    elements.append(Spacer(1, 8 * mm))

    # 四、AI 模块分析
    module_analyses = report_content.get('module_analyses', [])
    if module_analyses:
        elements.append(Paragraph('四、AI 模块分析', style_h1))
        for ma in module_analyses:
            elements.append(Paragraph(f"{ma['module_name']}（{ma['module_pct_score']:.2f}分）", style_h2))
            if ma.get('overall_evaluation'):
                elements.append(Paragraph(ma['overall_evaluation'], style_body))
            if ma.get('main_issues'):
                elements.append(Paragraph('<b>主要问题：</b>', style_bold))
                for iss in ma['main_issues']:
                    elements.append(Paragraph(f"• {iss}", style_bullet))
            if ma.get('improvement_suggestions'):
                elements.append(Paragraph('<b>改进建议：</b>', style_bold))
                for sug in ma['improvement_suggestions']:
                    elements.append(Paragraph(f"• {sug}", style_bullet))
            elements.append(Spacer(1, 4 * mm))

    # 五、综合分析报告
    ai_full_report = report_content.get('ai_full_report', '')
    if ai_full_report:
        elements.append(Paragraph('五、综合分析报告', style_h1))
        for line in ai_full_report.split('\n'):
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith('### '):
                elements.append(Paragraph(stripped[4:], style_h2))
            elif stripped.startswith('## '):
                elements.append(Paragraph(stripped[3:], style_h2))
            elif stripped.startswith('# '):
                elements.append(Paragraph(stripped[2:], style_h1))
            elif stripped.startswith('- ') or stripped.startswith('* '):
                elements.append(Paragraph(f"• {stripped[2:]}", style_bullet))
            elif stripped.startswith('|'):
                continue  # 跳过 Markdown 表格行（PDF 表格处理复杂）
            else:
                elements.append(Paragraph(stripped, style_body))

    doc.build(elements)
    return output_path


def generate_word_report(report_content: dict, output_path: str) -> str:
    """生成 Word 格式报告（增强版：包含AI分析内容）"""
    from docx import Document
    from docx.shared import Inches, Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()

    # 标题
    title = doc.add_heading('物业品质检查报告', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 一、检查概况
    doc.add_heading('一、检查概况', level=1)
    info_table = doc.add_table(rows=5, cols=2)
    info_table.style = 'Table Grid'
    cells = [
        ('项目名称', report_content.get('project_name', '')),
        ('检查日期', str(report_content.get('check_date', ''))),
        ('项目总分', f"{report_content.get('total_score', 0):.2f} 分"),
        ('问题总数', str(report_content.get('issue_summary', {}).get('total_count', 0))),
        ('评分方法', f"{report_content.get('scoring_method', '')} + {report_content.get('analysis_method', '')}")
    ]
    for i, (label, value) in enumerate(cells):
        info_table.rows[i].cells[0].text = label
        info_table.rows[i].cells[1].text = value
    doc.add_paragraph()

    # 二、各模块得分
    doc.add_heading('二、各模块检查情况', level=1)
    score_table = doc.add_table(rows=len(report_content.get('modules', [])) + 1, cols=5)
    score_table.style = 'Table Grid'
    headers = ['模块名称', '得分', '权重', '贡献分', '问题数']
    for i, header in enumerate(headers):
        score_table.rows[0].cells[i].text = header
    for i, module in enumerate(report_content.get('modules', []), 1):
        if module.get('module_pct_score', 0) > 0:
            score_table.rows[i].cells[0].text = module['module_name']
            score_table.rows[i].cells[1].text = f"{module['module_pct_score']:.2f}"
            score_table.rows[i].cells[2].text = f"{module['weight_ratio'] * 100:.0f}%"
            score_table.rows[i].cells[3].text = f"{module['weighted_contribution']:.2f}"
            score_table.rows[i].cells[4].text = str(module.get('issue_count', 0))
    doc.add_paragraph()

    # 三、重点问题清单
    doc.add_heading('三、重点问题清单', level=1)
    for severity in ['严重', '一般', '轻微']:
        issues = report_content.get('issues', {}).get(severity, [])
        if issues:
            idx = ['严重', '一般', '轻微'].index(severity) + 1
            doc.add_heading(f'3.{idx} {severity}问题（{len(issues)}项）', level=2)
            for issue in issues[:10]:
                p = doc.add_paragraph()
                p.add_run(f"[{issue.get('module_name')}] {issue.get('item_name', '')}").bold = True
                if issue.get('description'):
                    p.add_run(f"\n  问题：{issue['description']}")
                if issue.get('location'):
                    p.add_run(f"\n  位置：{issue['location']}")
    doc.add_paragraph()

    # 四、AI 模块分析（DeepSeek）
    module_analyses = report_content.get('module_analyses', [])
    if module_analyses:
        doc.add_heading('四、AI 模块分析', level=1)
        for ma in module_analyses:
            doc.add_heading(f"{ma['module_name']}（{ma['module_pct_score']:.2f}分）", level=2)
            if ma.get('overall_evaluation'):
                doc.add_paragraph(ma['overall_evaluation'])
            if ma.get('main_issues'):
                p = doc.add_paragraph()
                p.add_run('主要问题：').bold = True
                for issue_text in ma['main_issues']:
                    doc.add_paragraph(f"• {issue_text}", style='List Bullet')
            if ma.get('improvement_suggestions'):
                p = doc.add_paragraph()
                p.add_run('改进建议：').bold = True
                for suggestion in ma['improvement_suggestions']:
                    doc.add_paragraph(f"• {suggestion}", style='List Bullet')
        doc.add_paragraph()

    # 五、综合分析报告（DeepSeek 生成的 Markdown）
    ai_full_report = report_content.get('ai_full_report', '')
    if ai_full_report:
        doc.add_heading('五、综合分析报告', level=1)
        # 简易 Markdown → Word 转换
        for line in ai_full_report.split('\n'):
            line = line.strip()
            if not line:
                continue
            if line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line[2:], style='List Bullet')
            else:
                doc.add_paragraph(line)

    # 保存文档
    doc.save(output_path)
    return output_path


# ==================== API 端点 ====================
@router.post("/generate/{task_id}", summary="生成报告")
async def generate_report(
    task_id: str,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    current_user: User = Depends(check_role(["admin"]))
):
    """生成检查报告（Agent 3: DeepSeek-V3.2 AI分析）"""
    # 验证任务存在
    task = db.query(InspectionTask).filter(InspectionTask.task_id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")

    # 检查是否已评分
    results = db.query(ScoringResult).join(
        InspectionRecord, ScoringResult.record_id == InspectionRecord.record_id
    ).filter(InspectionRecord.task_id == task_id).first()
    if not results:
        raise HTTPException(status_code=400, detail="请先完成评分")

    # 检查是否正在生成
    progress = report_progress.get(task_id, {})
    if progress.get("status") == "generating":
        raise HTTPException(status_code=409, detail="报告正在生成中，请稍候")

    # 速率限制：60秒内不能重复触发
    now = time.time()
    if task_id in _report_trigger_times:
        elapsed = now - _report_trigger_times[task_id]
        if elapsed < _REPORT_RATE_LIMIT:
            raise HTTPException(
                status_code=429,
                detail=f"报告生成过于频繁，请{int(_REPORT_RATE_LIMIT - elapsed)}秒后再试"
            )
    _report_trigger_times[task_id] = now

    # 启动后台报告生成任务
    background_tasks.add_task(run_report_generation_sync, task_id)

    return {
        "message": "报告生成已启动",
        "task_id": task_id,
        "analysis_method": "DeepSeek-V3.2 AI分析"
    }


@router.get("/status/{task_id}", summary="查询报告生成进度")
async def get_report_status(
    task_id: str,
    current_user: User = Depends(get_current_user)
):
    """查询报告生成进度"""
    progress = report_progress.get(task_id)
    if not progress:
        # 内存没有 → 尝试从数据库恢复
        try:
            from core.progress_store import progress_store
            db_data = progress_store.get("report", task_id)
            if db_data:
                detail = db_data.get("detail", {})
                progress = {
                    "status": db_data["status"],
                    "current_step": detail.get("current_step", ""),
                    "modules_done": detail.get("modules_done", 0),
                    "modules_total": detail.get("modules_total", 0),
                    "error": db_data.get("error"),
                }
                # 回填到内存
                report_progress[task_id] = progress
        except Exception:
            pass
    if not progress:
        progress = {"status": "not_started"}
    return {
        "task_id": task_id,
        "status": progress.get("status", "not_started"),
        "current_step": progress.get("current_step", ""),
        "modules_done": progress.get("modules_done", 0),
        "modules_total": progress.get("modules_total", 0),
        "error": progress.get("error")
    }


@router.get("/list/all", summary="获取报告列表")
async def list_reports(
    page: int = 1,
    page_size: int = 20,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """获取报告列表（按项目隔离）"""
    project_filter = get_user_project_filter(current_user, db)
    if project_filter is not None and len(project_filter) == 0:
        return {"total": 0, "page": page, "page_size": page_size, "items": []}

    query = db.query(Report).join(InspectionTask).options(joinedload(Report.task).joinedload(InspectionTask.project))

    if project_filter is not None:
        query = query.filter(InspectionTask.project_id.in_(project_filter))

    total = query.count()
    reports = query.order_by(Report.generated_at.desc()).offset((page - 1) * page_size).limit(page_size).all()

    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": [
            {
                "report_id": r.report_id,
                "task_id": r.task_id,
                "project_name": r.task.project.name if r.task and r.task.project else "",
                "standard_type": r.task.standard_type if r.task else "diecheng",
                "total_score": float(r.total_score) if r.total_score else 0,
                "generated_at": r.generated_at.isoformat()
            }
            for r in reports
        ]
    }


@router.get("/{task_id}", summary="获取报告详情")
async def get_report(
    task_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """获取报告详情（按项目隔离）"""
    report = db.query(Report).filter(Report.task_id == task_id).order_by(Report.generated_at.desc()).first()
    if not report:
        raise HTTPException(status_code=404, detail="报告不存在")

    # 项目级隔离检查
    project_filter = get_user_project_filter(current_user, db)
    if project_filter is not None and len(project_filter) > 0:
        task = db.query(InspectionTask).filter(InspectionTask.task_id == task_id).first()
        if not task or task.project_id not in project_filter:
            raise HTTPException(status_code=403, detail="无权访问此报告")

    content = json.loads(report.content_json) if report.content_json else {}

    return {
        "report_id": report.report_id,
        "task_id": report.task_id,
        "project_name": content.get('project_name', ''),
        "check_date": content.get('check_date', ''),
        "total_score": float(report.total_score) if report.total_score else 0,
        "generated_at": report.generated_at.isoformat(),
        "content": content
    }


@router.get("/{task_id}/download", summary="下载报告文件")
async def download_report(
    task_id: str,
    request: Request,
    format: str = "word",
    token: str = None,
    db: Session = Depends(get_db)
):
    """下载报告文件（支持 word/pdf 格式）
    支持两种认证方式：Authorization header 或 token 查询参数（移动端下载用）
    """
    # 尝试从 header 认证
    current_user = None
    auth_header = request.headers.get("authorization", "")
    if auth_header.startswith("Bearer "):
        try:
            auth_token = auth_header.split(" ", 1)[1]
            payload = jwt.decode(auth_token, settings.SECRET_KEY, algorithms=[settings.ALGORITHM])
            user_id = payload.get("user_id")
            username = payload.get("sub")
            if user_id:
                current_user = db.query(User).filter(User.id == user_id).first()
            elif username:
                current_user = db.query(User).filter(User.username == username).first()
        except Exception:
            pass

    # header 认证失败，尝试 token 查询参数
    if current_user is None and token:
        try:
            payload = jwt.decode(token, settings.SECRET_KEY, algorithms=[settings.ALGORITHM])
            user_id = payload.get("user_id")
            username = payload.get("sub")
            if user_id:
                current_user = db.query(User).filter(User.id == user_id).first()
            elif username:
                current_user = db.query(User).filter(User.username == username).first()
        except Exception:
            pass

    if not current_user:
        raise HTTPException(status_code=401, detail="未认证")
    report = db.query(Report).filter(Report.task_id == task_id).order_by(Report.generated_at.desc()).first()
    if not report:
        raise HTTPException(status_code=404, detail="报告不存在")

    # 项目级隔离检查
    project_filter = get_user_project_filter(current_user, db)
    if project_filter is not None and len(project_filter) > 0:
        task = db.query(InspectionTask).filter(InspectionTask.task_id == task_id).first()
        if not task or task.project_id not in project_filter:
            raise HTTPException(status_code=403, detail="无权下载此报告")

    if format == "pdf":
        # PDF 路径：与 Word 同目录，扩展名为 .pdf
        word_path = report.file_path or ""
        pdf_path = word_path.replace('.docx', '.pdf')
        if not pdf_path or not os.path.exists(pdf_path):
            raise HTTPException(status_code=404, detail="PDF文件不存在，请重新生成报告")
        return FileResponse(
            pdf_path,
            media_type="application/pdf",
            filename=f"{report.report_id}.pdf"
        )
    else:
        if not report.file_path or not os.path.exists(report.file_path):
            raise HTTPException(status_code=404, detail="Word文件不存在")
        return FileResponse(
            report.file_path,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=f"{report.report_id}.docx"
        )


@router.get("/versions/{task_id}", summary="获取报告版本列表")
async def get_report_versions(
    task_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """获取某任务的所有报告版本（按版本号倒序）"""
    # 项目级隔离检查
    project_filter = get_user_project_filter(current_user, db)
    if project_filter is not None and len(project_filter) > 0:
        task = db.query(InspectionTask).filter(InspectionTask.task_id == task_id).first()
        if not task or task.project_id not in project_filter:
            raise HTTPException(status_code=403, detail="无权访问此报告")

    reports = db.query(Report).filter(
        Report.task_id == task_id
    ).order_by(Report.version.desc()).all()

    return {
        "task_id": task_id,
        "total_versions": len(reports),
        "versions": [
            {
                "report_id": r.report_id,
                "version": r.version,
                "total_score": float(r.total_score) if r.total_score else None,
                "generated_at": r.generated_at.isoformat() if r.generated_at else None,
                "file_exists": bool(r.file_path and os.path.exists(r.file_path))
            }
            for r in reports
        ]
    }


@router.get("/version/{report_id}", summary="获取指定版本报告详情")
async def get_report_by_id(
    report_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """通过 report_id 获取指定版本的报告详情"""
    report = db.query(Report).filter(Report.report_id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="报告不存在")

    # 项目级隔离检查
    project_filter = get_user_project_filter(current_user, db)
    if project_filter is not None and len(project_filter) > 0:
        task = db.query(InspectionTask).filter(InspectionTask.task_id == report.task_id).first()
        if not task or task.project_id not in project_filter:
            raise HTTPException(status_code=403, detail="无权访问此报告")

    content = json.loads(report.content_json) if report.content_json else {}

    return {
        "report_id": report.report_id,
        "task_id": report.task_id,
        "version": report.version,
        "project_name": content.get('project_name', ''),
        "check_date": content.get('check_date', ''),
        "total_score": float(report.total_score) if report.total_score else 0,
        "generated_at": report.generated_at.isoformat() if report.generated_at else None,
        "content": content
    }


@router.delete("/{task_id}", summary="删除报告")
async def delete_report(
    task_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(check_role(["admin"]))
):
    """删除报告及关联文件（仅管理员）"""
    report = db.query(Report).filter(Report.task_id == task_id).order_by(Report.generated_at.desc()).first()
    if not report:
        raise HTTPException(status_code=404, detail="报告不存在")

    report_id = report.report_id

    # 删除磁盘文件
    if report.file_path and os.path.exists(report.file_path):
        os.remove(report.file_path)
    pdf_path = report.file_path.replace('.docx', '.pdf') if report.file_path else ''
    if pdf_path and os.path.exists(pdf_path):
        os.remove(pdf_path)

    db.delete(report)
    db.commit()

    logger.info(f"报告已删除: {report_id} (task={task_id}), by user={current_user.username}")
    return {"message": "报告已删除", "report_id": report_id}

