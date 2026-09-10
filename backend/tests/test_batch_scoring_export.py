import io
import uuid

import openpyxl
from PIL import Image

from api.tasks import load_template_items
from core import standards
from models.models import (
    InspectionRecord,
    InspectionTask,
    Issue,
    Photo,
    ScoringResult,
    TaskAssignment,
)


def _uid(prefix):
    return f"{prefix}-{uuid.uuid4().hex[:12]}"


def _create_task(db, project, *, status="completed", standard_type="diecheng", total_score=80):
    task = InspectionTask(
        task_id=_uid("Q"),
        project_id=project.id,
        check_date="2026-08-11",
        status=status,
        standard_type=standard_type,
        total_score=total_score,
    )
    db.add(task)
    db.commit()
    db.refresh(task)
    return task


def _add_score(db, task, *, score=4, module_name=None, item_id="ITEM-001", item_name="检查项"):
    module_name = module_name or standards.get_modules(task.standard_type or "diecheng")[0]
    record = InspectionRecord(
        record_id=_uid("REC"),
        task_id=task.task_id,
        project_id=task.project_id,
        module_name=module_name,
        check_date=task.check_date,
        status="completed",
    )
    db.add(record)
    db.flush()
    result = ScoringResult(
        scoring_id=_uid("SCR"),
        record_id=record.record_id,
        module_name=module_name,
        item_id=item_id,
        item_name=item_name,
        score=score,
        weight=1,
        weighted_score=score,
        check_standard="检查要求",
        check_method="现场检查",
        scoring_rule="按检查标准评分",
        scoring_basis="AI评分依据",
        improvement_suggestion="整改建议",
        confidence_score=0.88,
        is_skipped=False,
    )
    db.add(result)
    db.commit()
    return record, result


def _open_workbook(response):
    assert response.status_code == 200, response.text
    return openpyxl.load_workbook(io.BytesIO(response.content))


def test_batch_selected_exports_three_sheets_photos_and_skips_unscored(
    client, db_session, admin_headers, test_project, tmp_path
):
    scored = _create_task(db_session, test_project, total_score=0)
    record, _ = _add_score(db_session, scored, score=0)
    unscored = _create_task(db_session, test_project, total_score=None)

    issue = Issue(
        issue_id=_uid("ISS"),
        record_id=record.record_id,
        module_name=record.module_name,
        item_id="ITEM-001",
        item_name="检查项",
        description="现场发现问题",
        location="一层",
        severity="一般",
    )
    db_session.add(issue)
    db_session.flush()
    image_path = tmp_path / "evidence.png"
    Image.new("RGB", (20, 10), color="red").save(image_path)
    db_session.add_all([
        Photo(photo_id=_uid("PHO"), issue_id=issue.issue_id, file_path=str(image_path), file_name="evidence.png"),
        Photo(photo_id=_uid("PHO"), issue_id=issue.issue_id, file_path=str(tmp_path / "missing.png"), file_name="missing.png"),
    ])
    db_session.commit()

    list_response = client.get("/api/v1/tasks/", headers=admin_headers)
    listed = {item["task_id"]: item for item in list_response.json()["items"]}
    assert listed[scored.task_id]["has_scoring_result"] is True
    assert listed[unscored.task_id]["has_scoring_result"] is False

    response = client.post(
        "/api/v1/scoring/export/batch",
        headers=admin_headers,
        json={"scope": "selected", "task_ids": [scored.task_id, unscored.task_id]},
    )
    assert response.headers["X-Exported-Count"] == "1"
    assert response.headers["X-Skipped-Count"] == "1"
    workbook = _open_workbook(response)
    assert workbook.sheetnames == ["任务汇总", "模块汇总", "检查项明细"]
    assert workbook["任务汇总"].max_row == 2
    assert workbook["任务汇总"]["G2"].value == 0
    detail_sheet = workbook["检查项明细"]
    assert detail_sheet.max_row == 2
    assert len(detail_sheet._images) == 1
    assert "图片缺失" in str(detail_sheet.cell(row=2, column=21).value)
    workbook.close()


def test_filtered_export_crosses_task_center_pagination_and_respects_filter(
    client, db_session, admin_headers, test_project
):
    for index in range(21):
        task = _create_task(db_session, test_project, status="completed", total_score=index)
        _add_score(db_session, task, score=5)
    pending = _create_task(db_session, test_project, status="pending", total_score=70)
    _add_score(db_session, pending, score=5)

    response = client.post(
        "/api/v1/scoring/export/batch",
        headers=admin_headers,
        json={"scope": "filtered", "project_id": test_project.id, "status": "completed"},
    )
    assert response.headers["X-Exported-Count"] == "21"
    assert response.headers["X-Skipped-Count"] == "0"
    workbook = _open_workbook(response)
    assert workbook["任务汇总"].max_row == 22
    exported_task_ids = {workbook["任务汇总"].cell(row=row, column=2).value for row in range(2, 23)}
    assert pending.task_id not in exported_task_ids
    workbook.close()


def test_batch_and_single_export_enforce_inspector_assignments(
    client, db_session, admin_headers, inspector_headers, inspector_user, test_project
):
    allowed = _create_task(db_session, test_project)
    denied = _create_task(db_session, test_project)
    _add_score(db_session, allowed)
    _add_score(db_session, denied)
    db_session.add(TaskAssignment(
        task_id=allowed.task_id,
        module_name=standards.get_modules("diecheng")[0],
        inspector_id=inspector_user.id,
    ))
    db_session.commit()

    denied_batch = client.post(
        "/api/v1/scoring/export/batch",
        headers=inspector_headers,
        json={"scope": "selected", "task_ids": [denied.task_id]},
    )
    assert denied_batch.status_code == 403
    denied_single = client.get(f"/api/v1/scoring/export/{denied.task_id}", headers=inspector_headers)
    assert denied_single.status_code == 403

    filtered = client.post(
        "/api/v1/scoring/export/batch",
        headers=inspector_headers,
        json={"scope": "filtered", "status": "completed"},
    )
    assert filtered.status_code == 200
    assert filtered.headers["X-Exported-Count"] == "1"


def test_lizhi_item_score_above_five_is_preserved(
    client, db_session, admin_headers, test_project
):
    module_name = standards.get_modules("lizhi")[0]
    template_items = load_template_items(module_name, "lizhi")
    item = next((row for row in template_items if float(row.get("max_score") or 0) > 5), template_items[0])
    maximum = float(item.get("max_score") or standards.get_module_max_score("lizhi", module_name))
    score = min(maximum, 10)
    task = _create_task(db_session, test_project, standard_type="lizhi", total_score=score)
    _add_score(
        db_session,
        task,
        score=score,
        module_name=module_name,
        item_id=item["item_id"],
        item_name=item.get("item_name") or "砺质检查项",
    )

    response = client.post(
        "/api/v1/scoring/export/batch",
        headers=admin_headers,
        json={"scope": "selected", "task_ids": [task.task_id]},
    )
    workbook = _open_workbook(response)
    detail_sheet = workbook["检查项明细"]
    assert detail_sheet.cell(row=2, column=11).value == maximum
    assert detail_sheet.cell(row=2, column=13).value == score
    assert detail_sheet.cell(row=2, column=13).value > 5
    workbook.close()
