"""
整改 API 测试
覆盖：整改列表获取、状态筛选、权限控制
"""
from io import BytesIO

from openpyxl import load_workbook
from PIL import Image

from models.models import InspectionRecord, InspectionTask, Issue, Photo, Rectification


class TestRectificationList:
    """整改记录列表测试"""

    def test_get_rectifications_without_auth(self, client):
        """无 token 访问整改列表应返回 401"""
        resp = client.get("/api/v1/rectifications/all")
        assert resp.status_code == 401

    def test_get_rectifications_with_auth(self, client, admin_headers):
        """有 token 应能获取整改列表"""
        resp = client.get("/api/v1/rectifications/all", headers=admin_headers)
        assert resp.status_code == 200
        data = resp.json()
        assert "items" in data or isinstance(data, list)

    def test_get_rectifications_with_status_filter(self, client, admin_headers):
        """状态筛选参数应生效"""
        for status_val in ["pending", "submitted", "approved"]:
            resp = client.get(f"/api/v1/rectifications/all?status={status_val}",
                              headers=admin_headers)
            assert resp.status_code == 200

    def test_get_rectifications_with_severity_filter(self, client, admin_headers):
        """严重程度筛选应生效"""
        for sev in ["严重", "一般", "轻微"]:
            resp = client.get(f"/api/v1/rectifications/all?severity={sev}",
                              headers=admin_headers)
            assert resp.status_code == 200

    def test_get_rectifications_pagination(self, client, admin_headers):
        """分页参数应生效"""
        resp = client.get("/api/v1/rectifications/all?page=1&page_size=10",
                          headers=admin_headers)
        assert resp.status_code == 200

    def test_get_pending_without_auth(self, client):
        """无 token 访问待整改列表应返回 401"""
        resp = client.get("/api/v1/rectifications/pending")
        assert resp.status_code == 401


class TestRectificationSubmit:
    """整改提交测试"""

    def test_submit_without_auth(self, client):
        """无 token 提交整改应返回 401"""
        resp = client.post("/api/v1/rectifications/nonexistent-id/submit", json={
            "rectification_note": "测试整改说明"
        })
        assert resp.status_code == 401

    def test_submit_nonexistent_rectification(self, client, admin_headers):
        """提交不存在的整改记录应返回 404"""
        resp = client.post("/api/v1/rectifications/99999/submit", json={
            "rectification_note": "测试整改说明"
        }, headers=admin_headers)
        assert resp.status_code in (404, 400)


class TestRectificationByRole:
    """不同角色的整改访问权限"""

    def test_inspector_can_list(self, client, inspector_headers):
        """检查员应能查看整改列表"""
        resp = client.get("/api/v1/rectifications/all", headers=inspector_headers)
        assert resp.status_code == 200


class TestRectificationExport:
    """整改导出测试"""

    def test_export_allows_shared_issue_photo_in_multiple_rows(
        self, client, admin_headers, db_session, test_project, tmp_path
    ):
        """同一问题对应多轮整改记录时，共用照片仍应能生成有效 Excel。"""
        task = InspectionTask(
            task_id="T-EXPORT-SHARED-PHOTO",
            project_id=test_project.id,
            check_date="2026-05-27",
        )
        record = InspectionRecord(
            record_id="REC-EXPORT-SHARED-PHOTO",
            task_id=task.task_id,
            project_id=test_project.id,
            module_name="客户服务",
            check_date="2026-05-27",
        )
        issue = Issue(
            issue_id="ISS-EXPORT-SHARED-PHOTO",
            record_id=record.record_id,
            module_name="客户服务",
            item_id="1.1",
            item_name="前台环境",
            description="台面杂乱",
        )
        photo_path = tmp_path / "shared.jpg"
        Image.new("RGB", (32, 24), (220, 30, 30)).save(photo_path, "JPEG")
        photo = Photo(
            photo_id="PHO-EXPORT-SHARED-PHOTO",
            issue_id=issue.issue_id,
            file_path=str(photo_path),
            file_name=photo_path.name,
            photo_type="问题照片",
        )
        rectifications = [
            Rectification(
                rectification_id=f"RECT-EXPORT-SHARED-{index}",
                issue_id=issue.issue_id,
                record_id=record.record_id,
                task_id=task.task_id,
            )
            for index in (1, 2)
        ]
        db_session.add_all([task, record, issue, photo, *rectifications])
        db_session.commit()

        response = client.get("/api/v1/rectifications/export/excel", headers=admin_headers)

        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook.active
        assert worksheet.max_row == 3
        assert len(worksheet._images) == 2
